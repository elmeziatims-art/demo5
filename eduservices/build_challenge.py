#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_challenge.py — le masque de challenge operationnel.

CE QU'IL DEMONTRE. Le CFO cadre au niveau groupe, le moteur descend sa cible a
la maille fine, et un directeur de campus corrige les seuls taux qu'il maitrise.
Le CA, les couts et l'EBITDA se recalculent, et l'ecart a la cible groupe se
referme -- ou pas -- sous les yeux de tout le monde.

POURQUOI LE MASQUE RECALCULE AU LIEU D'ATTENDRE TAGETIK. Un directeur qui
attend un aller-retour serveur entre chaque essai n'essaie pas. L'entonnoir est
donc rejoue dans la feuille, a l'identique de V_MOTEUR :

    leads rejoues = leads 2026 x [ org x (1+marque)^rendement_marque
                                 + payant x (1+acquisition)^rendement_acq ] / leads
    nouveaux      = leads rejoues x (conv_lead_cand + delta) x conv_cand_admis
                                  x (conv_admis_inscrit + delta)
    effectif      = nouveaux si annee d'entree, sinon effectif N-1 x (passage + delta)
    prix          = tarif 2026 x (1 + hausse x coefficient de marque)
    CA            = effectif x prix + nouveaux x frais de dossier

Cale contre la base : sur Pigier Bordeaux, cette reproduction rend
1 191 512,54 EUR contre 1 191 512,60 dans l'allocation -- six centimes.

L'EBITDA suit la regle de V_BUDGET et pas une autre : seuls les comptes 621,
604 et 6063 suivent le chiffre d'affaires. D'ou un taux marginal de 87,05 % --
chaque euro de CA gagne en laisse 0,87 en EBITDA -- et surtout un masque qui ne
peut pas diverger du budget que lit le CFO.
"""
import warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as GL
warnings.filterwarnings("ignore")

OUT = "/home/user/demo5/eduservices/CHALLENGE_OPERATIONNEL.xlsx"

#  la charte du masque Cadrage, pas une autre
INK, AZUR, GRIS = "262626", "007AC3", "E7E6E6"
CIEL, JAUNE, DOUX = "E8F1F9", "FFF2CC", "6B7075"
PANEL, FILET, ROUGE, VERT = "FFFFFF", "D5D7DA", "C41822", "2E7D4F"
UI = "Arial"
F = lambda sz=8.5, b=False, c=INK, i=False: Font(name=UI, size=sz, bold=b, color=c, italic=i)
fill = lambda c: PatternFill("solid", fgColor=c)
sd = lambda c=FILET, st="thin": Side(style=st, color=c)
D_ = Alignment("right", vertical="center")
C_ = Alignment("center", vertical="center")
ind = lambda n=0: Alignment("left", vertical="center", indent=n)

EUR  = '#,##0" €";[Red]-#,##0" €";"–"'
EURS = '+#,##0" €";[Red]-#,##0" €";"–"'
EUR2 = '#,##0.00" €"'
TAUX = '+0.00%;[Red]-0.00%;"–"'
PCT  = '0.0%;[Red]-0.0%;"–"'
NB   = '#,##0.0'
NBS  = '+#,##0.0;[Red]-#,##0.0;"–"'
COEF = '0.000000'

# ---------------------------------------------------------------- la donnee
PROPOSITION = [
    ("IPAC_MTP", "Ipac Bachelor Factory", "Montpellier", 122.132581, 818445.7303, 45651.4283, 138588.9216, 60270.6210, 285086.1631),
    ("IPAC_NAN", "Ipac Bachelor Factory", "Nantes", 151.322531, 1067200.8998, 74584.3246, 179282.7902, 79053.4471, 405388.6815),
    ("IPAC_REN", "Ipac Bachelor Factory", "Rennes", 122.121592, 818371.3392, 74584.3246, 145155.7596, 60952.5740, 284971.0992),
    ("ISCOM_LIL", "ISCOM", "Lille", 255.086642, 1844124.6190, 124306.1780, 303457.3352, 120718.7206, 690691.5518),
    ("ISCOM_PAR", "ISCOM", "Paris", 369.352343, 3048653.6062, 186459.7818, 520261.8117, 201288.8986, 1157580.0732),
    ("ISCOM_TLS", "ISCOM", "Toulouse", 241.481357, 1709754.9257, 124306.1780, 285351.0054, 112350.4412, 625795.4430),
    ("MBWAY_BOR", "MBway", "Bordeaux", 260.411810, 1863775.8159, 124306.1780, 298839.3379, 122314.1017, 689721.5139),
    ("MBWAY_LYO", "MBway", "Lyon", 338.498202, 2620797.0556, 161598.3403, 426605.5248, 169256.7809, 985233.8737),
    ("MBWAY_NAN", "MBway", "Nantes", 307.342957, 2267018.4682, 149167.6195, 366973.0584, 146693.0233, 837601.2482),
    ("MBWAY_PAR", "MBway", "Paris", 398.782457, 3292150.8481, 186459.7818, 535199.1154, 220105.8620, 1254602.3804),
    ("PIGIER_BOR", "Pigier", "Bordeaux", 184.716640, 1191512.6045, 99444.7365, 194818.2691, 78129.0202, 360701.1172),
    ("PIGIER_LYO", "Pigier", "Lyon", 236.947532, 1653870.0637, 99444.7365, 257235.0742, 109269.5504, 558386.8523),
    ("TUNON_LYO", "Tunon", "Lyon", 121.859881, 902399.2802, 74584.3246, 198220.2695, 65182.9340, 389149.6272),
    ("TUNON_PAR", "Tunon", "Paris", 143.560305, 1133628.5551, 74584.3246, 238111.8977, 85529.6614, 520675.9391),
]
CIBLE_GRP, PROP_GRP = 4095766.36, 4111503.64

CAMPUS = "PIGIER_BOR"
VILLE, MARQUE = "Bordeaux", "Pigier"
CAMPAGNE = dict(ORG=652, PAID=578, LEAD=1230, REND_ACQ=0.474597, REND_BRAND=0.357574)
COEF_MARQUE = 0.90
#  programme, annee, modalite, leads, conv lead>cand, conv cand>admis,
#  conv admis>inscrit, taux de passage, effectif N-1, tarif, classes, entree
LIGNES = [
    ("BTS Gestion", "1re année", "alternance", 615, .20, .69918699, .41860465, 0, 0, 5820., 2, 1, 30),
    ("BTS Gestion", "2e année", "alternance", 0, 0, 0, 0, .97222222, 36., 5820., 2, 0, 30),
    ("Bachelor RH", "1re année", "alternance", 615, .20, .69918699, .41860465, 0, 0, 6790., 2, 1, 32),
    ("Bachelor RH", "2e année", "alternance", 0, 0, 0, 0, .97222222, 36., 6790., 2, 0, 32),
    ("Bachelor RH", "3e année", "alternance", 0, 0, 0, 0, .97142857, 35., 6790., 2, 0, 32),
]
LEVIERS = [("Budget d'acquisition", "HYP_ACQ_BUD", .08, True),
           ("Budget de marque", "HYP_BRAND_BUD", .10, False),
           ("Hausse tarifaire", "HYP_PRICE", .0029, True),
           ("Conversion lead → candidature", "HYP_CNV_LEAD_CAND", .01, True),
           ("Conversion admis → inscrit", "HYP_CNV_ADM_INS", .01, True),
           ("Taux de passage", "HYP_PASS_RATE", .005, True)]
PRODUITS_26, VARIABLES_26 = 1132650., 149421.
PROD_LEVIER, FRAIS = .0185, 90.
#  la compta 2026 du campus, compte par compte : les charges NON variables se
#  projettent avec leurs propres leviers et ne bougent pas quand le campus
#  challenge. On les fige donc une fois, projetees.
COMPTES_26 = {"621": 96583., "604": 31703., "6063": 21135., "6231": 21968.,
              "6411": 182756.35, "6413": 73539.85, "645": 114395.03,
              "613": 112878.51, "615": 16125.25, "616": 8171.39, "625": 12256.21,
              "63511": 10750.17}
LEV_INFL, LEV_SAL, LEV_FTE, LEV_STRUCT = .02, .025, .04, .0
CHARGES_FIXES_27 = (
    COMPTES_26["6231"] * (1 + .08)
    + (COMPTES_26["6411"] + COMPTES_26["6413"] + COMPTES_26["645"])
      * (1 + LEV_SAL) * (1 + LEV_FTE)
    + (COMPTES_26["613"] + COMPTES_26["615"] + COMPTES_26["616"] + COMPTES_26["625"])
      * (1 + LEV_INFL) * (1 - PROD_LEVIER) * (1 + LEV_STRUCT)
    + COMPTES_26["63511"] * (1 + LEV_INFL) * (1 - PROD_LEVIER))

wb = openpyxl.Workbook()

def mettre(ws, r, c, v, font=None, al=None, fmt=None, fl=None, bd=None):
    x = ws.cell(r, c, v)
    x.font = font or F()
    x.alignment = al or ind(0)
    x.number_format = fmt or "General"
    if fl: x.fill = fill(fl)
    if bd: x.border = bd
    return x

def bandeau(ws, ncol, titre, sous):
    for c in range(1, ncol + 1):
        ws.cell(3, c).fill = fill(AZUR)
    mettre(ws, 2, 2, titre, F(14, True, INK), ind(0))
    mettre(ws, 2, 4, sous, F(9, False, DOUX, True), ind(0))
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 3

def section(ws, r, c1, c2, titre, note=""):
    mettre(ws, r, c1, titre, F(10.5, True, INK), ind(0))
    if note:
        mettre(ws, r, c2, note, F(7.5, False, DOUX, True), D_)
    for c in range(c1, c2 + 1):
        ws.cell(r, c).border = Border(bottom=sd(INK))
    ws.row_dimensions[r].height = 24

def entete(ws, r, c1, libelles, aligns, azur=()):
    for i, lib in enumerate(libelles):
        x = mettre(ws, r, c1 + i, lib,
                   F(8, True, "FFFFFF" if i in azur else INK),
                   {"g": ind(1), "d": D_, "c": C_}[aligns[i]],
                   fl=AZUR if i in azur else GRIS)
        x.border = Border(top=sd(INK), bottom=sd(INK))
    ws.row_dimensions[r].height = 20

def saisie(ws, r, c):
    x = ws.cell(r, c)
    x.fill = fill(JAUNE)
    x.border = Border(left=sd(INK), right=sd(INK), top=sd(INK), bottom=sd(INK))
    return x

def blanc(ws, nl, nc):
    for r in range(1, nl + 1):
        for c in range(1, nc + 1):
            ws.cell(r, c).fill = fill(PANEL)

# ==========================================================================
#  ONGLET 1 — LA PROPOSITION
# ==========================================================================
p1 = wb.active
p1.title = "Proposition"
p1.sheet_view.showGridLines = False
blanc(p1, 40, 10)
for c, w in ((1, 2.4), (2, 22), (3, 22), (4, 14), (5, 16), (6, 16), (7, 16), (8, 12), (9, 2.4)):
    p1.column_dimensions[GL(c)].width = w
bandeau(p1, 9, "① La proposition", "ce que le moteur descend de la cible groupe, campus par campus")
mettre(p1, 5, 2, "Cible EBITDA groupe", F(9, True), ind(0))
mettre(p1, 5, 4, CIBLE_GRP, F(11, True), D_, EUR)
mettre(p1, 5, 5, "Proposition", F(9, True), D_)
mettre(p1, 5, 6, PROP_GRP, F(11, True), D_, EUR)
mettre(p1, 5, 7, "Écart", F(9, True), D_)
mettre(p1, 5, 8, PROP_GRP - CIBLE_GRP, F(11, True, VERT), D_, EURS)
mettre(p1, 6, 2, "La proposition n'est pas une quote-part : c'est ce que produit la mécanique "
                 "de chaque campus — son entonnoir, ses cohortes, ses classes.",
       F(8, False, DOUX, True), ind(0))
section(p1, 8, 2, 8, "PAR CAMPUS", "scénario Cadrage · version V01")
entete(p1, 9, 2, ["Campus", "Marque", "Effectif", "Chiffre d'affaires",
                  "Coûts campus", "Marge campus", "Marge %"],
       ["g", "g", "d", "d", "d", "d", "d"])
for i, (ent, mq, ville, eff, ca, vac, perm, odir, struct) in enumerate(PROPOSITION):
    r = 10 + i
    couts = vac + perm + odir + struct
    p1.row_dimensions[r].height = 17
    mettre(p1, r, 2, ville, F(8.5, ent == CAMPUS), ind(1))
    mettre(p1, r, 3, mq, F(8.5, False, DOUX), ind(0))
    mettre(p1, r, 4, eff, F(8.5), D_, NB)
    mettre(p1, r, 5, ca, F(8.5, ent == CAMPUS), D_, EUR)
    mettre(p1, r, 6, couts, F(8.5), D_, EUR)
    mettre(p1, r, 7, ca - couts, F(8.5), D_, EUR)
    mettre(p1, r, 8, "=IF(E%d=0,0,G%d/E%d)" % (r, r, r), F(8.5), D_, PCT)
    for c in range(2, 9):
        p1.cell(r, c).border = Border(bottom=sd("EDEEF0"))
        if ent == CAMPUS:
            p1.cell(r, c).fill = fill(CIEL)
FIN1 = 10 + len(PROPOSITION)
p1.row_dimensions[FIN1].height = 19
mettre(p1, FIN1, 2, "Réseau", F(9, True), ind(1))
for c, f in ((4, "=SUM(D10:D%d)" % (FIN1 - 1)), (5, "=SUM(E10:E%d)" % (FIN1 - 1)),
             (6, "=SUM(F10:F%d)" % (FIN1 - 1)), (7, "=SUM(G10:G%d)" % (FIN1 - 1))):
    mettre(p1, FIN1, c, f, F(9, True), D_, NB if c == 4 else EUR)
mettre(p1, FIN1, 8, "=IF(E%d=0,0,G%d/E%d)" % (FIN1, FIN1, FIN1), F(9, True), D_, PCT)
for c in range(2, 9):
    p1.cell(FIN1, c).fill = fill(GRIS)
    p1.cell(FIN1, c).border = Border(top=sd(INK), bottom=Side(style="double", color=INK))
mettre(p1, FIN1 + 2, 2, "Source : V_MOTEUR et V_ALLOCATION, scénario 2027BUD_V1, version V01. "
                        "La ligne en bleu est le campus ouvert au challenge.",
       F(7.5, False, DOUX, True), ind(0))
p1.freeze_panes = "B10"

# ==========================================================================
#  ONGLET 2 — MON PERIMETRE  (le coeur : tout se recalcule ici)
# ==========================================================================
p2 = wb.create_sheet("Mon périmètre")
p2.sheet_view.showGridLines = False
blanc(p2, 70, 14)
for c, w in ((1, 2.4), (2, 24), (3, 13), (4, 13), (5, 13), (6, 13), (7, 13),
             (8, 13), (9, 15), (10, 15), (11, 2.4), (12, 20), (13, 13)):
    p2.column_dimensions[GL(c)].width = w
bandeau(p2, 11, "② Mon périmètre  ·  %s %s" % (MARQUE, VILLE),
        "j'ajuste ce que je maîtrise, le reste suit")

#  --- le bloc technique, a droite, masque : c'est la restitution ----------
mettre(p2, 5, 12, "Restitution", F(8, True, DOUX), ind(0))
TECH = [("Leads organiques 2026", CAMPAGNE["ORG"], NB),
        ("Leads payants 2026", CAMPAGNE["PAID"], NB),
        ("Leads totaux 2026", CAMPAGNE["LEAD"], NB),
        ("Rendement acquisition", CAMPAGNE["REND_ACQ"], COEF),
        ("Rendement marque", CAMPAGNE["REND_BRAND"], COEF),
        ("Coefficient prix marque", COEF_MARQUE, '0.00'),
        ("Frais de dossier", FRAIS, EUR),
        ("Produits 2026", PRODUITS_26, EUR),
        ("Variables 2026  621/604/6063", VARIABLES_26, EUR),
        ("Charges fixes projetées 2027", CHARGES_FIXES_27, EUR),
        ("Effort de productivité", PROD_LEVIER, TAUX)]
for i, (lib, val, fmt) in enumerate(TECH):
    r = 6 + i
    mettre(p2, r, 12, lib, F(8, False, DOUX), ind(0))
    mettre(p2, r, 13, val, F(8, False, DOUX), D_, fmt)
T = {lib: "$M$%d" % (6 + i) for i, (lib, v, f) in enumerate(TECH)}
for c in (12, 13):
    p2.column_dimensions[GL(c)].hidden = True

#  --- (a) les leviers ------------------------------------------------------
section(p2, 5, 2, 5, "ⓐ   LES LEVIERS DE MON PÉRIMÈTRE",
        "jaune = je saisis · gris = décidé au-dessus de moi")
entete(p2, 6, 2, ["Levier", "Proposé", "Ajusté", "Écart"],
       ["g", "c", "c", "c"], azur=(2,))
LIG0 = 7
for i, (lib, code, val, ouvert) in enumerate(LEVIERS):
    r = LIG0 + i
    p2.row_dimensions[r].height = 18
    mettre(p2, r, 2, lib, F(8.5), ind(1))
    mettre(p2, r, 3, val, F(8.5, False, DOUX), C_, TAUX)
    if ouvert:
        mettre(p2, r, 4, val, F(9, True), C_, TAUX)
        saisie(p2, r, 4)
    else:
        mettre(p2, r, 4, "décidé par la marque", F(7.5, False, DOUX, True), C_, fl=GRIS)
    mettre(p2, r, 5, "=IF(ISNUMBER(D%d),D%d-C%d,0)" % (r, r, r), F(8.5, False, VERT), C_, TAUX)
    for c in range(2, 6):
        p2.cell(r, c).border = Border(bottom=sd("EDEEF0"))
ACQ, BRAND, PRIX_L, GLC, GCV, PASS = ("$D$%d" % (LIG0 + i) for i in range(6))
BRAND = "$C$%d" % (LIG0 + 1)            # la marque n'est pas ajustable ici
FIN_L = LIG0 + len(LEVIERS)

#  --- le facteur de reponse des leads -------------------------------------
mettre(p2, FIN_L + 1, 2, "Effet sur mes leads", F(8.5, True), ind(1))
FACT = ("=({o}*(1+{b})^{rb}+{p}*(1+{a})^{ra})/{l}"
        .format(o=T["Leads organiques 2026"], b=BRAND, rb=T["Rendement marque"],
                p=T["Leads payants 2026"], a=ACQ, ra=T["Rendement acquisition"],
                l=T["Leads totaux 2026"]))
mettre(p2, FIN_L + 1, 4, FACT, F(9, True), C_, COEF)
p2.cell(FIN_L + 1, 4).fill = fill(CIEL)
mettre(p2, FIN_L + 1, 5, "rendement %.3f — doubler la dépense ne double pas les leads   "
                         ""
       % CAMPAGNE["REND_ACQ"], F(7.5, False, DOUX, True), ind(0))
FA = "$D$%d" % (FIN_L + 1)

#  --- (b) le recalcul, ligne a ligne --------------------------------------
R0 = FIN_L + 3
section(p2, R0, 2, 10, "ⓑ   CE QUE ÇA DONNE, LIGNE À LIGNE",
        "l'entonnoir rejoué à l'identique du moteur")
entete(p2, R0 + 1, 2, ["Programme", "Année", "Leads", "Nouveaux", "Effectif",
                       "Prix", "CA ajusté", "CA proposé", "Écart"],
       ["g", "g", "d", "d", "d", "d", "d", "d", "d"], azur=(6,))
L0 = R0 + 2
for i, (pg, an, mo, lead, rlc, rca, yld, pas, effinf, rev, cls, entree, capa) in enumerate(LIGNES):
    r = L0 + i
    p2.row_dimensions[r].height = 17
    mettre(p2, r, 2, pg, F(8.5), ind(1))
    mettre(p2, r, 3, an, F(8.5, False, DOUX), ind(0))
    mettre(p2, r, 4, "=%d*%s" % (lead, FA), F(8.5), D_, NB)
    if entree:
        mettre(p2, r, 5, "=D{r}*({rlc}+{glc})*{rca}*({yld}+{gcv})"
               .format(r=r, rlc=rlc, glc=GLC, rca=rca, yld=yld, gcv=GCV), F(8.5), D_, NB)
        mettre(p2, r, 6, "=E%d" % r, F(8.5), D_, NB)
    else:
        mettre(p2, r, 5, 0, F(8.5, False, DOUX), D_, NB)
        mettre(p2, r, 6, "={ei}*({p}+{pass_})".format(ei=effinf, p=pas, pass_=PASS),
               F(8.5), D_, NB)
    mettre(p2, r, 7, "={rev}*(1+{pr}*{co})".format(rev=rev, pr=PRIX_L, co=T["Coefficient prix marque"]),
           F(8.5), D_, EUR2)
    mettre(p2, r, 8, "=F{r}*G{r}+E{r}*{fe}".format(r=r, fe=T["Frais de dossier"]),
           F(8.5, True), D_, EUR)
    p2.cell(r, 8).fill = fill(CIEL)
    mettre(p2, r, 9, 0, F(8.5, False, DOUX), D_, EUR)      # rempli plus bas
    mettre(p2, r, 10, "=H%d-I%d" % (r, r), F(8.5), D_, EURS)
    for c in range(2, 11):
        p2.cell(r, c).border = Border(bottom=sd("EDEEF0"))
LN = L0 + len(LIGNES)
p2.row_dimensions[LN].height = 19
mettre(p2, LN, 2, "Total campus", F(9, True), ind(1))
for c in (5, 6):
    mettre(p2, LN, c, "=SUM(%s%d:%s%d)" % (GL(c), L0, GL(c), LN - 1), F(9, True), D_, NB)
for c in (8, 9, 10):
    mettre(p2, LN, c, "=SUM(%s%d:%s%d)" % (GL(c), L0, GL(c), LN - 1), F(9, True), D_,
           EURS if c == 10 else EUR)
for c in range(2, 11):
    p2.cell(LN, c).fill = fill(GRIS)
    p2.cell(LN, c).border = Border(top=sd(INK), bottom=Side(style="double", color=INK))

#  --- la colonne « CA proposé » : le moteur avec les leviers du groupe -----
#  On la calcule ici plutot que de la coller en dur : si un levier groupe
#  change, la reference change avec lui.
def moteur(lev):
    f = ((CAMPAGNE["ORG"] * (1 + lev["BRAND"]) ** CAMPAGNE["REND_BRAND"]
        + CAMPAGNE["PAID"] * (1 + lev["ACQ"]) ** CAMPAGNE["REND_ACQ"]) / CAMPAGNE["LEAD"])
    out, ca, eff, ent_ = [], 0.0, 0.0, 0.0
    for pg, an, mo, lead, rlc, rca, yld, pas, effinf, rev, cls, entree, capa in LIGNES:
        nc = lead * f * (rlc + lev["GLC"]) * rca * (yld + lev["GCV"])
        n = nc if entree else 0.0
        e = nc if entree else effinf * (pas + lev["PASS"])
        prix = rev * (1 + lev["PRICE"] * COEF_MARQUE)
        c = e * prix + n * FRAIS
        out.append(c); ca += c; eff += e; ent_ += n
    return out, ca, eff, ent_

LEV_GRP = dict(ACQ=LEVIERS[0][2], BRAND=LEVIERS[1][2], PRICE=LEVIERS[2][2],
               GLC=LEVIERS[3][2], GCV=LEVIERS[4][2], PASS=LEVIERS[5][2])
CA_LIGNES, CA_PROP, EFF_PROP, ENTREE_PROP = moteur(LEV_GRP)
for i, v in enumerate(CA_LIGNES):
    p2.cell(L0 + i, 9).value = v

#  --- (c) ce que ça change ------------------------------------------------
E0 = LN + 2
section(p2, E0, 2, 5, "ⓒ   CE QUE ÇA CHANGE",
        "règle de coûts identique à V_BUDGET : seuls 621, 604 et 6063 suivent le CA")
entete(p2, E0 + 1, 2, ["", "Proposé", "Ajusté", "Écart"],
       ["g", "d", "d", "d"], azur=(2,))
#  seuls les comptes variables suivent le CA : chaque euro de CA gagne laisse
#  donc (1 - part variable) en EBITDA. C'est exactement la regle de V_BUDGET,
#  et c'est ce qui empeche le masque de diverger du budget que lit le CFO.
MARGINAL = "(1-{v}*(1-{pr})/{p})".format(v=T["Variables 2026  621/604/6063"],
                                         pr=T["Effort de productivité"], p=T["Produits 2026"])
EB_PROP = ("={ca}-({v}*({ca}/{p})*(1-{pr})+{cf})"
           .format(ca="$I$%d" % LN, v=T["Variables 2026  621/604/6063"],
                   p=T["Produits 2026"], pr=T["Effort de productivité"],
                   cf=T["Charges fixes projetées 2027"]))
BLOC = [("Chiffre d'affaires", "=$I$%d" % LN, "=$H$%d" % LN, EUR),
        ("Effectif", "%.6f" % EFF_PROP, "=$F$%d" % LN, NB),
        ("EBITDA du campus", EB_PROP, None, EUR)]
for i, (lib, prop, aju, fmt) in enumerate(BLOC):
    r = E0 + 2 + i
    p2.row_dimensions[r].height = 18
    gras = lib.startswith("EBITDA")
    mettre(p2, r, 2, lib, F(9, gras), ind(1))
    mettre(p2, r, 3, prop if str(prop).startswith("=") else float(prop),
           F(9, gras), D_, fmt)
    if aju is None:
        aju = "=C{r}+(H{ln}-I{ln})*{m}".format(r=r, ln=LN, m=MARGINAL)
    mettre(p2, r, 4, aju, F(9, gras), D_, fmt)
    mettre(p2, r, 5, "=D%d-C%d" % (r, r), F(9, gras, VERT), D_,
           EURS if fmt == EUR else NBS)
    for c in range(2, 6):
        p2.cell(r, c).border = Border(bottom=sd("EDEEF0"))
        if gras:
            p2.cell(r, c).fill = fill(CIEL)
            p2.cell(r, c).border = Border(top=sd(AZUR), bottom=sd(AZUR))
R_EB = E0 + 4
mettre(p2, E0 + 6, 2, "Taux marginal", F(8.5), ind(1))
mettre(p2, E0 + 6, 3, "=" + MARGINAL, F(8.5, True), D_, PCT)
mettre(p2, E0 + 6, 5, "part de chaque euro de CA supplémentaire qui tombe en EBITDA — "
                      "les autres comptes ne suivent pas le chiffre d'affaires",
       F(7.5, False, DOUX, True), ind(0))

#  --- (d) le garde-fou de capacite ----------------------------------------
#  Il se teste sur les COHORTES D'ENTREE, pas sur le campus : les places de
#  2e et 3e annee existent mais on n'y entre pas.
K0 = E0 + 8
PLACES = sum(cls * capa for pg, an, mo, ld, a, b, c_, d, e_, rv, cls, ent, capa in LIGNES if ent)
section(p2, K0, 2, 5, "ⓓ   AI-JE LA PLACE ?",
        "testé sur les cohortes d'entrée, jamais sur le campus entier")
ENTREES = "+".join("E%d" % (L0 + i) for i, l in enumerate(LIGNES) if l[11])
CAP = [("Places en cohorte d'entrée", str(PLACES), NB),
       ("Nouveaux inscrits attendus", "=" + ENTREES, NB),
       ("Taux de remplissage à l'entrée", "=IF(%d=0,0,(%s)/%d)" % (PLACES, ENTREES, PLACES), PCT),
       ("Classes à ouvrir", "=MAX(0,ROUNDUP(((%s)-%d)/%d,0))"
        % (ENTREES, PLACES, LIGNES[0][12]), '#,##0')]
for i, (lib, val, fmt) in enumerate(CAP):
    r = K0 + 1 + i
    p2.row_dimensions[r].height = 17
    mettre(p2, r, 2, lib, F(8.5), ind(1))
    mettre(p2, r, 3, val if str(val).startswith("=") else float(val), F(9, i >= 2), D_, fmt)
    for c in range(2, 6):
        p2.cell(r, c).border = Border(bottom=sd("EDEEF0"))
p2.conditional_formatting.add("C%d" % (K0 + 3), CellIsRule(
    operator="greaterThan", formula=["1"], font=Font(name=UI, size=9, bold=True, color=ROUGE)))
p2.conditional_formatting.add("C%d" % (K0 + 4), CellIsRule(
    operator="greaterThan", formula=["0"], font=Font(name=UI, size=9, bold=True, color=ROUGE)))
mettre(p2, K0 + 6, 2, "Le réseau a 974 places libres mais seulement 351 en première année : "
                      "les autres sont en 2e et 3e année, on n'y entre pas. Un contrôle au "
                      "niveau du campus laisserait passer des inscriptions qu'aucune classe "
                      "ne peut prendre.", F(7.5, False, DOUX, True), ind(0))
p2.freeze_panes = "B7"
CA_AJ, EB_AJ = "'Mon périmètre'!$H$%d" % LN, "'Mon périmètre'!$D$%d" % R_EB
CA_PR, EB_PR = "'Mon périmètre'!$I$%d" % LN, "'Mon périmètre'!$C$%d" % R_EB

# ==========================================================================
#  ONGLET 3 — CIBLE VS CONSTRUIT
# ==========================================================================
p3 = wb.create_sheet("Cible vs Construit")
p3.sheet_view.showGridLines = False
blanc(p3, 40, 10)
for c, w in ((1, 2.4), (2, 40), (3, 18), (4, 18), (5, 16), (6, 2.4)):
    p3.column_dimensions[GL(c)].width = w
bandeau(p3, 6, "③ Cible vs construit", "où en est le groupe une fois les campus passés")
section(p3, 5, 2, 5, "LE PONT", "scénario Cadrage · version V01")
entete(p3, 6, 2, ["Étape", "EBITDA", "Cumul", "Écart à la cible"],
       ["g", "d", "d", "d"], azur=(1,))
PONT = [("Cible du CFO", CIBLE_GRP, "socle"),
        ("Proposition du moteur, avant challenge", PROP_GRP, "niveau"),
        ("Effet %s %s" % (MARQUE, VILLE), None, "effet"),
        ("Budget construit", None, "total")]
r = 7
for lib, val, genre in PONT:
    p3.row_dimensions[r].height = 19
    gras = genre in ("socle", "total")
    mettre(p3, r, 2, lib, F(9, gras), ind(1))
    if genre == "socle":
        mettre(p3, r, 3, val, F(9, True), D_, EUR)
        mettre(p3, r, 4, "=C%d" % r, F(9, True), D_, EUR)
        mettre(p3, r, 5, 0, F(9), D_, EURS)
        p3.cell(r, 2).fill = fill(GRIS); p3.cell(r, 3).fill = fill(GRIS)
        p3.cell(r, 4).fill = fill(GRIS); p3.cell(r, 5).fill = fill(GRIS)
    elif genre == "niveau":
        mettre(p3, r, 3, val, F(9), D_, EUR)
        mettre(p3, r, 4, "=C%d" % r, F(9), D_, EUR)
        mettre(p3, r, 5, "=D%d-$C$7" % r, F(9), D_, EURS)
    elif genre == "effet":
        mettre(p3, r, 3, "=%s-%s" % (EB_AJ, EB_PR), F(9, False, VERT), D_, EURS)
        mettre(p3, r, 4, "=D%d+C%d" % (r - 1, r), F(9), D_, EUR)
        mettre(p3, r, 5, "=D%d-$C$7" % r, F(9), D_, EURS)
    else:
        mettre(p3, r, 3, "=D%d" % (r - 1), F(10, True), D_, EUR)
        mettre(p3, r, 4, "=C%d" % r, F(10, True), D_, EUR)
        mettre(p3, r, 5, "=D%d-$C$7" % r, F(10, True), D_, EURS)
        for c in range(2, 6):
            p3.cell(r, c).fill = fill(CIEL)
            p3.cell(r, c).border = Border(top=sd(AZUR),
                                          bottom=Side(style="double", color=INK))
    for c in range(2, 6):
        if p3.cell(r, c).border.bottom is None:
            p3.cell(r, c).border = Border(bottom=sd("EDEEF0"))
    r += 1
R_FIN = r - 1
p3.conditional_formatting.add("E7:E%d" % R_FIN, CellIsRule(
    operator="lessThan", formula=["0"], font=Font(name=UI, size=9, bold=True, color=ROUGE)))

mettre(p3, R_FIN + 2, 2, '="Treize campus sur quatorze n\'ont rien changé : ils héritent du '
                         'cadrage groupe. Le seul écart vient de %s, pour "'
                         '&TEXT(ABS(C%d),"#,##0 €")&"."' % (VILLE, R_FIN - 1),
       F(9, True, INK), ind(0))
p3.row_dimensions[R_FIN + 2].height = 20
for c in range(2, 6):
    p3.cell(R_FIN + 2, c).fill = fill(CIEL)

section(p3, R_FIN + 4, 2, 5, "QUI A CHALLENGÉ", "lu dans le cube : ENTITY, USERUPD, DATEUPD")
entete(p3, R_FIN + 5, 2, ["Campus", "Leviers repris", "Saisi par", "Le"],
       ["g", "c", "g", "g"])
r = R_FIN + 6
for ent, mq, ville, eff, ca, v_, pe, od, st in PROPOSITION:
    p3.row_dimensions[r].height = 16
    mettre(p3, r, 2, "%s %s" % (mq, ville), F(8.5, ent == CAMPUS), ind(1))
    if ent == CAMPUS:
        mettre(p3, r, 3, '=COUNTIF(\'Mon périmètre\'!$E$%d:$E$%d,"<>0")' % (LIG0, FIN_L - 1),
               F(8.5, True), C_, '#,##0')
        mettre(p3, r, 4, "directeur de campus", F(8.5), ind(0))
        mettre(p3, r, 5, "en séance", F(8.5, False, DOUX, True), ind(0))
        for c in range(2, 6):
            p3.cell(r, c).fill = fill(CIEL)
    else:
        mettre(p3, r, 3, 0, F(8.5, False, DOUX), C_, '#,##0')
        mettre(p3, r, 4, "—", F(8.5, False, DOUX), ind(0))
    for c in range(2, 6):
        p3.cell(r, c).border = Border(bottom=sd("EDEEF0"))
    r += 1
mettre(p3, r + 1, 2, "Aucune ligne n'est écrite ailleurs que dans AW_002_000001_000001. "
                     "Un campus qui ne challenge pas n'a aucune ligne : il hérite du groupe "
                     "par la cascade.", F(7.5, False, DOUX, True), ind(0))

# ==========================================================================
#  CONTROLES
# ==========================================================================
for f in (p1, p2, p3):
    for row in f.iter_rows():
        for cel in row:
            if cel.font is None or cel.font.name != UI:
                cel.font = F(8.5)
            if not cel.number_format:
                cel.number_format = "General"
wb.save(OUT)
v = openpyxl.load_workbook(OUT)
pb = []
for f in v.worksheets:
    if f.merged_cells.ranges:
        pb.append("%s : fusion" % f.title)
    for row in f.iter_rows():
        for cel in row:
            if isinstance(cel.value, str) and "_xlfn" in cel.value:
                pb.append("%s!%s : _xlfn" % (f.title, cel.coordinate))
            if cel.number_format == "":
                pb.append("%s!%s : format vide" % (f.title, cel.coordinate))
print("onglets :", v.sheetnames)
print("CA proposé du campus : %.2f   effectif %.3f   entrée %.3f / %d places"
      % (CA_PROP, EFF_PROP, ENTREE_PROP, PLACES))
print("fusions :", sum(len(f.merged_cells.ranges) for f in v.worksheets))
print("contrôle :", "PASS" if not pb else pb[:5])

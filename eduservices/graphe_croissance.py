# -*- coding: utf-8 -*-
"""graphe_croissance.py — le graphe de decomposition de la croissance du CA.

CE QU'IL MONTRE. Le bloc CROISSANCE aligne six leviers : cinq agissent sur les
volumes, un seul sur le prix. Le graphe pose la question en face : sur la
croissance du chiffre d'affaires, combien vient du volume, combien du prix ?

  effet volume = (q1 - q0) * p0        effet prix = (p1 - p0) * q1
  et leur somme vaut exactement q1*p1 - q0*p0. Verifie sur les trois scenarios,
  ecart 0,00 EUR.

POURQUOI PAS UNE CASCADE COMPLETE. Avec les niveaux de CA (23 M et 24 M) dans
le meme graphe, l'axe monte a 24 millions et l'effet prix -- 97 000 EUR --
mesure quatre dixiemes de pour cent de la hauteur : invisible. On ne trace donc
que les effets. Les niveaux sont deja dans les cartes du haut.

COMMENT IL EST GREFFE. On genere le XML du graphe avec openpyxl dans un
classeur jetable, puis on le greffe : le classeur de l'utilisateur porte un
graphique qu'openpyxl ne sait pas relire, un aller-retour le supprimerait.
"""
import re, sys, zipfile, openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (RichTextProperties, Paragraph, ParagraphProperties,
                                   CharacterProperties, Font as PoliceDessin)

SRC, DST = sys.argv[1], sys.argv[2]
CADRAGE = "xl/worksheets/sheet2.xml"
UI, NOIR = "Arial", "000000"
AZUR, AMBRE, ENCRE, GRIS = "007AC3", "B26B00", "1F3B57", "595959"

#  ------------------------------------------------------------------------
#  1. le graphe, fabrique par openpyxl dans un classeur jetable
#  ------------------------------------------------------------------------
def texte(taille=800, gras=False):
    cp = CharacterProperties(latin=PoliceDessin(typeface=UI), sz=taille, b=gras,
                             solidFill=NOIR)
    return RichText(bodyPr=RichTextProperties(),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])

tmp = openpyxl.Workbook()
ws = tmp.active
ws.title = "Cadrage"                      # les references doivent viser SA feuille
ws["T22"], ws["U22"] = "Étape", "Effet"
for i, lib in enumerate(("Effet volume", "Effet prix", "Croissance totale")):
    ws.cell(23 + i, 20, lib)
    ws.cell(23 + i, 21, 0)

g = BarChart()
g.type = "col"
g.grouping = "clustered"
g.gapWidth = 80
g.add_data(Reference(ws, min_col=21, min_row=22, max_row=25), titles_from_data=True)
g.set_categories(Reference(ws, min_col=20, min_row=23, max_row=25))
serie = g.series[0]
serie.graphicalProperties.solidFill = AZUR
points = []
for idx, coul in ((1, AMBRE), (2, ENCRE)):
    pt = DataPoint(idx=idx)
    pt.graphicalProperties.solidFill = coul
    points.append(pt)
serie.data_points = points
g.dLbls = DataLabelList()
g.dLbls.showVal = True
g.dLbls.numFmt = '#,##0," k€"'
g.dLbls.txPr = texte(750)
g.title = "D'où vient la croissance du chiffre d'affaires"
gras = CharacterProperties(latin=PoliceDessin(typeface=UI), sz=1000, b=True, solidFill=NOIR)
g.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=gras)
g.title.tx.rich.p[0].r[0].rPr = gras
g.y_axis.numFmt = '#,##0," k€"'
g.x_axis.delete = False
g.y_axis.delete = False
g.x_axis.axPos = "b"
g.y_axis.axPos = "l"
g.legend = None
g.txPr = texte(800)
ws.add_chart(g, "A1")
tmp.save("_graphe_tmp.xlsx")
chart_xml = zipfile.ZipFile("_graphe_tmp.xlsx").read("xl/charts/chart1.xml").decode("utf8")

#  ------------------------------------------------------------------------
#  2. la zone de donnees, en T21:U25 de Cadrage
#  ------------------------------------------------------------------------
#  q0 = C15 effectif 2026, q1 = D15 effectif construit, C12 et D12 les CA.
VOLUME = "($D$15-$C$15)*($C$12/$C$15)"
PRIX   = "($D$12/$D$15-$C$12/$C$15)*$D$15"
DONNEES = [("T21", "s", "Données du graphique"),
           ("T22", "s", "Étape"),          ("U22", "s", "Effet"),
           ("T23", "s", "Effet volume"),   ("U23", "f", VOLUME),
           ("T24", "s", "Effet prix"),     ("U24", "f", PRIX),
           ("T25", "s", "Croissance totale"), ("U25", "f", "$D$12-$C$12")]

def bloc(xml, conteneur):
    m = re.search(r"<%s\b[^>]*>.*?</%s>" % (conteneur, conteneur), xml, re.S)
    return (m.group(0), m.start(), m.end()) if m else (None, -1, -1)

def elements(b, enfant):
    return re.findall(r"<%s\b(?:[^>]*/>|[^>]*>.*?</%s>)" % (enfant, enfant), b, re.S)

def ajouter(xml, conteneur, enfant, neuf):
    b, d, f = bloc(xml, conteneur)
    liste = elements(b, enfant)
    if neuf in liste:
        return xml, liste.index(neuf)
    idx = len(liste)
    corps = b[: -len("</%s>" % conteneur)] + neuf + "</%s>" % conteneur
    corps = re.sub(r'count="\d+"', 'count="%d"' % (idx + 1), corps, count=1)
    return xml[:d] + corps + xml[f:], idx

zin = zipfile.ZipFile(SRC)
styles = zin.read("xl/styles.xml").decode("utf8")
feuille = zin.read(CADRAGE).decode("utf8")

styles, police = ajouter(styles, "fonts", "font",
    '<font><sz val="8"/><color rgb="FF%s"/><name val="%s"/></font>' % (GRIS, UI))
styles, style_note = ajouter(styles, "cellXfs", "xf",
    '<xf numFmtId="3" fontId="%d" fillId="0" borderId="0" applyNumberFormat="1" '
    'applyFont="1"/>' % police)

def ligne_de(feuille, n):
    return re.search(r'<row r="%d"[^>]*>.*?</row>' % n, feuille, re.S)

ech = lambda t: t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
pose = 0
for ref, genre, val in DONNEES:
    n = int(re.search(r"\d+", ref).group(0))
    interieur = ("<is><t>%s</t></is>" % ech(val)) if genre == "s" else ("<f>%s</f>" % ech(val))
    attrs = ' s="%d"%s' % (style_note, ' t="inlineStr"' if genre == "s" else "")
    cellule = '<c r="%s"%s>%s</c>' % (ref, attrs, interieur)
    m = re.search(r'<c r="%s"((?:\s+[a-zA-Z:]+="[^"]*")*)\s*(?:/>|>.*?</c>)' % ref,
                  feuille, re.S)
    if m:                                            # la cellule existe deja
        feuille = feuille[:m.start()] + cellule + feuille[m.end():]
    else:                                            # sinon on l'insere dans sa ligne
        r = ligne_de(feuille, n)
        assert r, "ligne %d absente de Cadrage" % n
        contenu = r.group(0)
        neuf = contenu[: contenu.rindex("</row>")] + cellule + "</row>"
        neuf = re.sub(r'(<row r="%d"[^>]*?)\s*spans="[^"]*"' % n, r"\1", neuf)
        feuille = feuille[:r.start()] + neuf + feuille[r.end():]
    pose += 1

#  ------------------------------------------------------------------------
#  3. la greffe : dessin, relations, types de contenu
#  ------------------------------------------------------------------------
dessin = zin.read("xl/drawings/drawing1.xml").decode("utf8")
rels = zin.read("xl/drawings/_rels/drawing1.xml.rels").decode("utf8")
types = zin.read("[Content_Types].xml").decode("utf8")

rid = "rId%d" % (max(int(i) for i in re.findall(r'Id="rId(\d+)"', rels)) + 1)
ident = max(int(i) for i in re.findall(r'<xdr:cNvPr id="(\d+)"', dessin)) + 1
rels = rels.replace("</Relationships>",
    '<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/'
    '2006/relationships/chart" Target="../charts/chart2.xml"/></Relationships>' % rid)
types = types.replace("</Types>",
    '<Override PartName="/xl/charts/chart2.xml" ContentType="application/vnd.'
    'openxmlformats-officedocument.drawingml.chart+xml"/></Types>')

ANCRE = (
 '<xdr:twoCellAnchor>'
 '<xdr:from><xdr:col>12</xdr:col><xdr:colOff>0</xdr:colOff>'
 '<xdr:row>20</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
 '<xdr:to><xdr:col>18</xdr:col><xdr:colOff>0</xdr:colOff>'
 '<xdr:row>38</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
 '<xdr:graphicFrame macro="">'
 '<xdr:nvGraphicFramePr><xdr:cNvPr id="%d" name="GrapheCroissance"/>'
 '<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>'
 '<xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>'
 '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
 '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
 'r:id="%s"/></a:graphicData></a:graphic></xdr:graphicFrame>'
 '<xdr:clientData/></xdr:twoCellAnchor>') % (ident, rid)
assert "</xdr:wsDr>" in dessin
dessin = dessin.replace("</xdr:wsDr>", ANCRE + "</xdr:wsDr>")

wbx = zin.read("xl/workbook.xml").decode("utf8")
wbx = re.sub(r"<calcPr\b[^>]*/>", '<calcPr calcId="191029" fullCalcOnLoad="1"/>', wbx) \
    if "<calcPr" in wbx else \
    wbx.replace("</workbook>", '<calcPr calcId="191029" fullCalcOnLoad="1"/></workbook>')

zout = zipfile.ZipFile(DST, "w", zipfile.ZIP_DEFLATED)
for n in zin.namelist():
    if n == "xl/calcChain.xml":
        continue
    d = zin.read(n)
    if n == "xl/styles.xml":                        d = styles.encode("utf8")
    elif n == CADRAGE:                              d = feuille.encode("utf8")
    elif n == "xl/drawings/drawing1.xml":           d = dessin.encode("utf8")
    elif n == "xl/drawings/_rels/drawing1.xml.rels": d = rels.encode("utf8")
    elif n == "[Content_Types].xml":                d = types.encode("utf8")
    elif n == "xl/workbook.xml":                    d = wbx.encode("utf8")
    zout.writestr(n, d)
zout.writestr("xl/charts/chart2.xml", chart_xml.encode("utf8"))
zout.close()
print("cellules posees :", pose, "| relation :", rid, "| id de forme :", ident)

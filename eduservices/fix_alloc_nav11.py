# -*- coding: utf-8 -*-
"""Correctif chirurgical : on n'ouvre pas le classeur avec openpyxl pour le
reecrire -- un aller-retour lui coutait vingt-et-un fichiers, dont
xl/webextensions (le volet du complement Excel) et customXml. On patche le XML.
"""
import re, io, zipfile, shutil, collections, warnings
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.formula.translate import Translator

SRC, DST = "src.xlsx", "ALLOC_nav11_corrige.xlsx"
DEB, GRP, FIN_OLD, FIN_NEW = 18, 97, 349, 420
S_ALLOC, S_CALC = "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"

#  --- 1. lire la matiere avec openpyxl (lecture seule, sur la copie nettoyee)
vv = openpyxl.load_workbook("clean.xlsx", data_only=True)
ff = openpyxl.load_workbook("clean.xlsx")
v, a, c = vv["Allocation"], ff["Allocation"], ff["_CALC_ALLOC"]

REST, ENT = set(), collections.defaultdict(set)
for r in range(2, 400):
    e = v.cell(r, 62).value
    if e is None:
        continue
    REST.add((e, v.cell(r, 64).value, v.cell(r, 65).value, v.cell(r, 66).value))
    ENT[v.cell(r, 63).value].add(e)

MARQUE = {"MBway": "MBWAY", "ISCOM": "ISCOM", "Ipac Bachelor Factory": "IPAC",
          "Pigier": "PIGIER", "Tunon": "TUNON"}
VILLE = {"Paris": "PAR", "Lyon": "LYO", "Nantes": "NAN", "Bordeaux": "BOR",
         "Lille": "LIL", "Toulouse": "TLS", "Rennes": "REN", "Montpellier": "MTP"}
PROG = {"Bachelor Commerce": "BAC_CCE", "Bachelor Communication": "BAC_COM",
        "Bachelor Management": "BAC_MGT", "Bachelor Ressources humaines": "BAC_RH",
        "Bachelor Tourisme": "BAC_TOU", "BTS Gestion": "BTS_GES",
        "Mastère Communication": "MAS_COM", "Mastère Management": "MAS_MGT"}
ANS = {"BAC": {"1re année": "B1", "2e année": "B2", "3e année": "B3"},
       "MAS": {"1re année": "M1", "2e année": "M2"},
       "BTS": {"1re année": "BTS1", "2e année": "BTS2"}}
MOD = {"initial": "INIT", "alternance": "ALT"}

def niveau(r):
    f = a.cell(r, 3).value or ""
    if "$V%d" % r in f: return 1
    if "$W%d" % r in f: return 3
    if "$U%d" % r in f: return 2
    return 0

CLES, mq, cp, vus = {}, None, None, []
for r in range(DEB, GRP):
    lib = (a.cell(r, 2).value or "").strip()
    niv = niveau(r)
    if niv == 1:
        mq = MARQUE[lib]; cp = None
        vals = (mq, "", mq, "", "", "")
    elif niv == 2:
        cp = "%s_%s" % (mq, VILLE[lib])
        assert cp in ENT[mq], cp
        vals = (cp, cp, mq, "", "", "")
    else:
        pl, al, ml = [t.strip() for t in lib.split(",")]
        p = PROG[pl]; an = ANS[p.split("_")[0]][al]; mo = MOD[ml]
        assert (cp, p, an, mo) in REST, str((cp, p, an, mo))
        vus.append((cp, p, an, mo))
        vals = ("|".join((cp, p, an, mo)), cp, mq, p, an, mo)
    CLES[r] = (vals, niv)
CLES[GRP] = (("GROUPE", "", "", "", "", ""), 0)
assert len(vus) == len(set(vus)) == len(REST) == 60, (len(vus), len(set(vus)), len(REST))
n = collections.Counter(niv for _, niv in CLES.values())
print("plan reconstitue : %d marques, %d campus, %d classes, %d total  |  "
      "les 60 classes couvertes une fois chacune" % (n[1], n[2], n[3], n[0]))

#  --- 2. les formules de la ligne 349, resolues puis translatees ------------
MODELE = [(c.cell(FIN_OLD, col).column_letter, c.cell(FIN_OLD, col).value)
          for col in range(1, 47) if c.cell(FIN_OLD, col).value is not None]

#  --- 3. le patch XML -------------------------------------------------------
def esc(t):
    return (str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

zin = zipfile.ZipFile(SRC)
tmp = DST + ".tmp"
zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)

for it in zin.infolist():
    data = zin.read(it.filename)

    if it.filename == S_ALLOC:
        s = data.decode("utf-8")
        #  les colonnes techniques : creees et masquees
        if 'min="20" max="26"' not in s:
            s = s.replace("<cols>", '<cols><col min="20" max="26" width="13" '
                                    'hidden="1" customWidth="1"/>', 1)
        for r, (vals, niv) in CLES.items():
            m = re.search(r'<row r="%d"[^>]*>' % r, s)
            assert m, r
            bal = m.group(0)
            bal = re.sub(r'\s+outlineLevel="\d+"', "", bal)
            bal = re.sub(r'\s+hidden="\d"', "", bal)
            attrs = ' outlineLevel="%d"' % max(niv - 1, 0) if niv else ""
            if niv == 3:
                attrs += ' hidden="1"'
            bal = bal[:-1] + attrs + ">"
            s = s[:m.start()] + bal + s[m.end():]
            #  les cellules T..Z, inserees juste avant la zone de restitution
            cells = "".join(
                '<c r="%s%d" t="inlineStr"><is><t>%s</t></is></c>' % (L, r, esc(x))
                for L, x in zip("TUVWXY", vals) if x != "") \
                + '<c r="Z%d"><v>%d</v></c>' % (r, niv)
            s = re.sub(r'(<c r="BF%d")' % r, cells + r"\1", s, count=1)
        #  le temoin : la cellule n'existait pas, il faut l'inserer apres C12
        tem = ('<c r="D12"><f>"version "&amp;$P$1&amp;"  \u00b7  "&amp;TEXT(COUNTIF('
               '_CALC_ALLOC!$AS$1:$AS$%d,"2027|"&amp;$P$1),"0")&amp;'
               '" classes lues sur 60"</f></c>' % FIN_NEW)
        s = re.sub(r'<c r="D12"[^>]*(?:/>|>.*?</c>)', "", s, count=1, flags=re.S)
        m12 = re.search(r'(<c r="C12"(?:[^>]*?/>|[^>]*?>.*?</c>))', s, re.S)
        assert m12, "C12 introuvable"
        s = s[:m12.end()] + tem + s[m12.end():]
        data = s.encode("utf-8")

    elif it.filename == S_CALC:
        s = data.decode("utf-8")
        blocs = []
        for r in range(FIN_OLD + 1, FIN_NEW + 1):
            cs = "".join('<c r="%s%d"><f>%s</f></c>'
                         % (L, r, esc(Translator(f, origin="%s%d" % (L, FIN_OLD))
                                      .translate_formula("%s%d" % (L, r))[1:]))
                         for L, f in MODELE if isinstance(f, str) and f.startswith("="))
            blocs.append('<row r="%d" spans="1:45">%s</row>' % (r, cs))
        s = s.replace("</sheetData>", "".join(blocs) + "</sheetData>")
        s = s.replace('<dimension ref="A1:AT%d"/>' % FIN_OLD,
                      '<dimension ref="A1:AT%d"/>' % FIN_NEW)
        data = s.encode("utf-8")

    elif it.filename == "xl/workbook.xml":
        s = data.decode("utf-8")
        s = re.sub(r'<calcPr[^/]*/>', '<calcPr calcId="0" fullCalcOnLoad="1"/>', s)
        data = s.encode("utf-8")

    elif it.filename == "xl/calcChain.xml":
        pass                          # on le garde : aucune cellule supprimee

    zout.writestr(it, data)

zout.close(); zin.close()
shutil.move(tmp, DST)
print("ecrit :", DST)

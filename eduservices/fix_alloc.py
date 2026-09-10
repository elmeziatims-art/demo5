# -*- coding: utf-8 -*-
"""fix_alloc.py — le tableau d'allocation reconstruit depuis la restitution.

DEUX CORRECTIFS, ET ILS SE TIENNENT.

1. DEUX LIGNES MANQUAIENT. Le tableau affichait 58 classes pour 60 restituees :
   BAC_RH B2 ALT absent sur les deux campus Pigier, 572 803,83 EUR. La cause
   est que la liste des lignes etait figee dans la feuille. On ne la fige plus :
   elle est ENGENDREE depuis la zone de restitution, donc elle ne peut plus
   perdre une classe sans qu'on perde la donnee elle-meme.

2. LES CODES SORTENT DE LA COLONNE LUE. La colonne B ne porte plus que du
   francais ; les codes passent dans un bloc technique T:Z, masque, ou chaque
   ligne porte sa cle composee et ses quatre composantes. Les SUMIFS ne citent
   plus de litteraux : ils lisent ces colonnes. Ajouter un campus ou un
   programme devient une ligne de plus, pas une formule a reecrire.

Au passage : IFERROR remplace par IF -- Tagetik le reprefixe en _xlfn a chaque
aller-retour -- et l'ordre du plan remis a summaryBelow=False, puisque la ligne
de synthese est AU-DESSUS de son groupe et non en dessous.
"""
import re, zipfile
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as GL

SRC, DST = "_a.xlsx", "ALLOC_CORRIGE.xlsx"
ORIG = "AL.xlsx"

MARQUES = {"MBWAY": "MBway", "ISCOM": "ISCOM", "IPAC": "Ipac Bachelor Factory",
           "PIGIER": "Pigier", "TUNON": "Tunon"}
PROGRAMMES = {"BAC_CCE": "Bachelor Commerce", "BAC_COM": "Bachelor Communication",
              "BAC_MGT": "Bachelor Management", "BAC_RH": "Bachelor Ressources humaines",
              "BAC_TOU": "Bachelor Tourisme", "BTS_GES": "BTS Gestion",
              "MAS_COM": "Mastère Communication", "MAS_MGT": "Mastère Management"}
ANNEES = {"B1": "1re année", "B2": "2e année", "B3": "3e année",
          "M1": "1re année", "M2": "2e année", "BTS1": "1re année", "BTS2": "2e année"}
MODALITES = {"INIT": "initial", "ALT": "alternance"}
RANG_AN = {"B1": 1, "B2": 2, "B3": 3, "M1": 1, "M2": 2, "BTS1": 1, "BTS2": 2}

wb = openpyxl.load_workbook(SRC)
vv = openpyxl.load_workbook(SRC, data_only=True)
s, v = wb["Allocation"], vv["Allocation"]

# ---- la matiere : la zone de restitution, pas la liste figee ---------------
#  BJ=62 ENTITY, BK=63 MARQUE, BL=64 PROGRAMME, BM=65 AN_ETUDE, BN=66 MODALITE
lignes, villes = [], {}
for r in range(2, 400):
    ent = v.cell(r, 62).value
    if ent is None:
        continue
    lignes.append((v.cell(r, 63).value, ent, v.cell(r, 64).value,
                   v.cell(r, 65).value, v.cell(r, 66).value))
lignes = sorted(set(lignes))
for r in range(18, 96):
    b = s.cell(r, 2).value
    if b and b.startswith("   ") and not b.startswith("      ") and "(" in b:
        villes[b.split("(")[1].rstrip(")")] = b.split("(")[0].strip()

#  l'ordre affiche : celui du tableau d'origine pour les marques et les campus,
#  puis programme et annee pour les classes.
ordre_mq, ordre_cp = [], []
for r in range(18, 96):
    b = s.cell(r, 2).value
    if not b or b.strip() == "GROUPE":
        continue
    if b.startswith("      "):
        continue
    if b.startswith("   "):
        ordre_cp.append(b.split("(")[1].rstrip(")"))
    else:
        ordre_mq.append(b.strip().upper().replace("IPAC BACHELOR FACTORY", "IPAC"))
plan = []
for mq in ordre_mq:
    plan.append((1, mq, None, None, None, None))
    for cp in [c for c in ordre_cp if c.startswith(mq + "_")]:
        plan.append((2, mq, cp, None, None, None))
        cls = sorted([l for l in lignes if l[1] == cp],
                     key=lambda l: (l[2], RANG_AN.get(l[3], 9)))
        for _, _, pg, an, mo in cls:
            plan.append((3, mq, cp, pg, an, mo))
assert len(plan) == len(ordre_mq) + len(ordre_cp) + len(lignes), "plan incomplet"

# ---- les styles, repris sur les lignes existantes --------------------------
GAB = {1: 18, 2: 19, 3: 20, 0: 95}
modele = {niv: [copy(s.cell(r, c)._style) for c in range(1, 14)] for niv, r in GAB.items()}
HAUT = s.row_dimensions[19].height

# ---- on reecrit tout le bloc ----------------------------------------------
DEB = 18
FIN_ANC = 95
for r in range(DEB, FIN_ANC + 1):
    for c in range(1, 27):
        s.cell(r, c).value = None

CAL = "_CALC_ALLOC!"
CLE = '{c}$AS$1:$AS$349,"2027|"&$P$1'.format(c=CAL)
MES = [(3, "$G$1:$G$349"), (4, "$I$1:$I$349"), (5, "$AI$1:$AI$349"),
       (6, "$AJ$1:$AJ$349"), (7, "$AK$1:$AK$349"), (8, "$AL$1:$AL$349"),
       (9, "$AR$1:$AR$349"), (10, "$AM$1:$AM$349"), (12, "$AN$1:$AN$349")]

def critere(niv, r):
    if niv == 1:
        return "{c}$C$1:$C$349,$V{r}".format(c=CAL, r=r)
    if niv == 2:
        return "{c}$B$1:$B$349,$U{r}".format(c=CAL, r=r)
    return ("{c}$B$1:$B$349,$U{r},{c}$D$1:$D$349,$W{r},"
            "{c}$E$1:$E$349,$X{r},{c}$F$1:$F$349,$Y{r}").format(c=CAL, r=r)

for i, (niv, mq, cp, pg, an, mo) in enumerate(plan):
    r = DEB + i
    s.row_dimensions[r].height = HAUT
    s.row_dimensions[r].outlineLevel = niv - 1
    for c in range(1, 14):
        s.cell(r, c)._style = copy(modele[niv][c - 1])
    if niv == 1:   lib = MARQUES[mq]
    elif niv == 2: lib = villes.get(cp, cp)
    else:          lib = "%s, %s, %s" % (PROGRAMMES[pg], ANNEES[an], MODALITES[mo])
    #  le retrait se fait par l'ALIGNEMENT, pas par des espaces en tete : des
    #  espaces se perdent au premier tri et faussent toute recherche exacte.
    x = s.cell(r, 2, lib)
    x.alignment = Alignment(horizontal="left", vertical="center", indent=niv - 1)
    crit = critere(niv, r)
    for col, plage in MES:
        s.cell(r, col, "=SUMIFS({c}{p},{k},{q})".format(c=CAL, p=plage, k=CLE, q=crit))
    s.cell(r, 11, "=E{r}+F{r}+G{r}+H{r}+I{r}+J{r}".format(r=r))
    s.cell(r, 13, "=IF(D{r}=0,0,L{r}/D{r})".format(r=r))
    #  le bloc technique : la cle composee d'abord, ses composantes ensuite
    cle = "|".join(x for x in (cp or mq, pg, an, mo) if x)
    for col, val in ((20, cle), (21, cp or ""), (22, mq), (23, pg or ""),
                     (24, an or ""), (25, mo or ""), (26, niv)):
        x = s.cell(r, col, val)
        x.font = Font(name="Arial", size=8, color="FF808080")

GRP = DEB + len(plan)
s.row_dimensions[GRP].height = HAUT
s.row_dimensions[GRP].outlineLevel = 0
for c in range(1, 14):
    s.cell(GRP, c)._style = copy(modele[0][c - 1])
x = s.cell(GRP, 2, "GROUPE")
x.alignment = Alignment(horizontal="left", vertical="center", indent=0)
for col, plage in MES:
    s.cell(GRP, col, "=SUMIFS({c}{p},{k})".format(c=CAL, p=plage, k=CLE))
s.cell(GRP, 11, "=E{r}+F{r}+G{r}+H{r}+I{r}+J{r}".format(r=GRP))
s.cell(GRP, 13, "=IF(D{r}=0,0,L{r}/D{r})".format(r=GRP))
for col, val in ((20, "GROUPE"), (22, ""), (26, 0)):
    x = s.cell(GRP, col, val); x.font = Font(name="Arial", size=8, color="FF808080")

# ---- l'en-tete du bloc technique ------------------------------------------
for col, lib in ((20, "Clé"), (21, "Entity"), (22, "Marque"), (23, "Programme"),
                 (24, "Année"), (25, "Modalité"), (26, "Niveau")):
    x = s.cell(17, col, lib)
    x.font = Font(name="Arial", size=8, bold=True, color="FF262626")
    x.fill = PatternFill("solid", fgColor="FFE7E6E6")
    x.alignment = Alignment("left", vertical="center")
    x.border = Border(top=Side(style="thin", color="FF262626"),
                      bottom=Side(style="thin", color="FF262626"))
for col in range(20, 27):
    s.column_dimensions[GL(col)].width = 16
    s.column_dimensions[GL(col)].hidden = True

# ---- le reste : plan, mise en forme conditionnelle -------------------------
s.sheet_properties.outlinePr.summaryBelow = False    # la synthese est AU-DESSUS
s.conditional_formatting._cf_rules.clear()
s.conditional_formatting.add("M%d:M%d" % (DEB, GRP), ColorScaleRule(
    start_type="min", start_color="F6C9CC", mid_type="percentile", mid_value=50,
    mid_color="FFFFFF", end_type="max", end_color="DCEBC0"))

wb.save(DST)

# ---- on rend au fichier ce que le passage par openpyxl lui prend -----------
zo = zipfile.ZipFile(DST); parts = {n: zo.read(n) for n in zo.namelist()}; zo.close()
zi = zipfile.ZipFile(ORIG)
wbx = parts["xl/workbook.xml"].decode("utf8")
for nom in ("ALLOC_BRAND_CAMP", "ALLOC_CAMP_CLASS", "ALLOC_GRP_BRAND", "ALLOC_GRP_MARQUE"):
    wbx = wbx.replace('<definedName name="%s">' % nom,
                      '<definedName name="%s" isReadOnly="false">' % nom)
parts["xl/workbook.xml"] = wbx.encode("utf8")
parts["docProps/custom.xml"] = zi.read("docProps/custom.xml")
types = parts["[Content_Types].xml"].decode("utf8")
if "custom.xml" not in types:
    types = types.replace("</Types>", '<Override PartName="/docProps/custom.xml" ContentType='
                          '"application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
    parts["[Content_Types].xml"] = types.encode("utf8")
rels = parts["_rels/.rels"].decode("utf8")
if "custom.xml" not in rels:
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdCustom" Type="http://schemas.'
                        'openxmlformats.org/officeDocument/2006/relationships/custom-properties"'
                        ' Target="docProps/custom.xml"/></Relationships>')
    parts["_rels/.rels"] = rels.encode("utf8")
#  SUMIFS et IFERROR sont des fonctions de 2007 : elles ne doivent JAMAIS
#  porter le prefixe _xlfn, que Tagetik reinjecte a chaque aller-retour. Tant
#  qu'il est la, la feuille de calcul rend #NOM? des qu'on ouvre hors Tagetik,
#  et c'est elle qui alimente tout l'onglet visible.
n_pref = 0
for nom in list(parts):
    if nom.startswith("xl/worksheets/sheet"):
        x = parts[nom].decode("utf8")
        neuf, k = re.subn(r"_xlfn\.(?=SUMIFS|IFERROR)", "", x)
        if k:
            parts[nom] = neuf.encode("utf8"); n_pref += k
wbx2 = parts["xl/workbook.xml"].decode("utf8")
parts["xl/workbook.xml"] = (
    re.sub(r"<calcPr\b[^>]*/>", '<calcPr calcId="191029" fullCalcOnLoad="1"/>', wbx2)
    if "<calcPr" in wbx2 else
    wbx2.replace("</workbook>", '<calcPr calcId="191029" fullCalcOnLoad="1"/></workbook>')
).encode("utf8")
parts.pop("xl/calcChain.xml", None)

zn = zipfile.ZipFile(DST, "w", zipfile.ZIP_DEFLATED)
for n, d in parts.items():
    zn.writestr(n, d)
zn.close()

print("lignes du plan : %d marques, %d campus, %d classes  ->  %d lignes + GROUPE en %d"
      % (len(ordre_mq), len(ordre_cp), len(lignes), len(plan), GRP))
print("isReadOnly restaure :", wbx.count('isReadOnly'))
print("prefixes _xlfn retires :", n_pref)

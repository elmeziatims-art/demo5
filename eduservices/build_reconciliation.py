#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_reconciliation.py — le masque de reconciliation CRM / Comptabilite.

LA CONTRAINTE QUI COMMANDE TOUT LE FICHIER : le rapport est dynamique, on ne
sait pas ou la restitution va sortir. Aucune formule du masque ne vise donc une
ADRESSE. Elles se reperent par le NOM des colonnes :

  MATCH("CA CRM", <ligne d'en-tete>, 0)  donne le numero de colonne
  INDEX(Donnees!$A:$Z, 0, ce numero)     rend la colonne entiere
  SUMIFS(...) somme dessus

Consequence : la restitution peut sortir avec autant de lignes qu'elle veut,
dans n'importe quel ordre de colonnes, le masque suit. La seule chose qu'il
faut lui dire est le NUMERO DE LA LIGNE D'EN-TETE, en C7. Si elle bouge, on
change ce chiffre et rien d'autre.

LE FILET DE SECURITE : chaque bloc porte une ligne "Hors liste" qui vaut le
total de la restitution moins la somme des lignes affichees. Si une marque ou
un campus apparait que le masque ne connait pas, il ne disparait pas en
silence -- il se voit.

Pas de IFERROR : Tagetik reinjecte des prefixes _xlfn sur cette fonction a
chaque aller-retour. On utilise IF et ISERROR, qui passent.
"""
import warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as GL
warnings.filterwarnings("ignore")

OUT = "/home/user/demo5/eduservices/RECONCILIATION_CRM_COMPTA.xlsx"

NOIR, ENCRE, TETE = "000000", "1F3B57", "D9E2EC"
BLEU, JAUNE, GRIS = "DDEBF7", "FFF2CC", "595959"
FIN, ROUGE, VERT = "BFBFBF", "C00000", "2E7D4F"
UI = "Arial"

F = lambda sz=9, b=False, c=NOIR, i=False: Font(name=UI, size=sz, bold=b, color=c, italic=i)
fill = lambda c: PatternFill("solid", fgColor=c)
sd = lambda c=FIN, st="thin": Side(style=st, color=c)
G_ = Alignment("left", vertical="center")
D_ = Alignment("right", vertical="center")
C_ = Alignment("center", vertical="center")
ind = lambda n=1: Alignment("left", vertical="center", indent=n)

EUR = '#,##0" \u20ac";[Red]-#,##0" \u20ac"'
EURS = '+#,##0" \u20ac";[Red]-#,##0" \u20ac";"\u2013"'
PCT = '+0.00%;[Red]-0.00%;"\u2013"'

def mettre(ws, r, c, v, font=None, al=None, fmt=None, fl=None, bd=None):
    x = ws.cell(r, c, v)
    x.font = font or F()
    x.alignment = al or G_
    x.number_format = fmt or "General"
    if fl: x.fill = fill(fl)
    if bd: x.border = bd
    return x

DONNEES_LIGNES = [
    ("IPAC_MTP", "2024", "ACT", 701490.00, 701490.00),
    ("IPAC_NAN", "2024", "ACT", 906960.00, 906960.00),
    ("IPAC_REN", "2024", "ACT", 701490.00, 701490.00),
    ("ISCOM_LIL", "2024", "ACT", 1561220.00, 1561220.00),
    ("ISCOM_PAR", "2024", "ACT", 2580240.00, 2580240.00),
    ("ISCOM_TLS", "2024", "ACT", 1451880.00, 1451880.00),
    ("MBWAY_BOR", "2024", "ACT", 1573680.00, 1573680.00),
    ("MBWAY_LYO", "2024", "ACT", 2227170.00, 2227170.00),
    ("MBWAY_NAN", "2024", "ACT", 1923040.00, 1923040.00),
    ("MBWAY_PAR", "2024", "ACT", 2793380.00, 2793380.00),
    ("PIGIER_BOR", "2024", "ACT", 1003890.00, 1003890.00),
    ("PIGIER_LYO", "2024", "ACT", 1399680.00, 1399680.00),
    ("TUNON_LYO", "2024", "ACT", 774990.00, 774990.00),
    ("TUNON_PAR", "2024", "ACT", 968100.00, 968100.00),
    ("IPAC_MTP", "2025", "ACT", 741570.00, 741570.00),
    ("IPAC_NAN", "2025", "ACT", 956230.00, 956230.00),
    ("IPAC_REN", "2025", "ACT", 741570.00, 741570.00),
    ("ISCOM_LIL", "2025", "ACT", 1648000.00, 1648000.00),
    ("ISCOM_PAR", "2025", "ACT", 2744480.00, 2744480.00),
    ("ISCOM_TLS", "2025", "ACT", 1522320.00, 1522320.00),
    ("MBWAY_BOR", "2025", "ACT", 1666855.00, 1666855.00),
    ("MBWAY_LYO", "2025", "ACT", 2342685.00, 2342685.00),
    ("MBWAY_NAN", "2025", "ACT", 2048260.00, 2048260.00),
    ("MBWAY_PAR", "2025", "ACT", 2974600.00, 2974600.00),
    ("PIGIER_BOR", "2025", "ACT", 1068270.00, 1068270.00),
    ("PIGIER_LYO", "2025", "ACT", 1469340.00, 1469340.00),
    ("TUNON_LYO", "2025", "ACT", 819270.00, 819270.00),
    ("TUNON_PAR", "2025", "ACT", 1015320.00, 1015320.00),
    ("IPAC_MTP", "2026", "ACT", 781650.00, 781650.00),
    ("IPAC_NAN", "2026", "ACT", 1019500.00, 1019500.00),
    ("IPAC_REN", "2026", "ACT", 781650.00, 781650.00),
    ("ISCOM_LIL", "2026", "ACT", 1755850.00, 1755850.00),
    ("ISCOM_PAR", "2026", "ACT", 2908720.00, 2908720.00),
    ("ISCOM_TLS", "2026", "ACT", 1627980.00, 1627980.00),
    ("MBWAY_BOR", "2026", "ACT", 1773610.00, 1773610.00),
    ("MBWAY_LYO", "2026", "ACT", 2496705.00, 2496705.00),
    ("MBWAY_NAN", "2026", "ACT", 2158300.00, 2158300.00),
    ("MBWAY_PAR", "2026", "ACT", 3138840.00, 3138840.00),
    ("PIGIER_BOR", "2026", "ACT", 1132650.00, 1132650.00),
    ("PIGIER_LYO", "2026", "ACT", 1573830.00, 1573830.00),
    ("TUNON_LYO", "2026", "ACT", 863550.00, 863550.00),
    ("TUNON_PAR", "2026", "ACT", 1086150.00, 1086150.00),
    ("IPAC_MTP", "2027", "V01", 818445.73, 819979.77),
    ("IPAC_NAN", "2027", "V01", 1067200.90, 1069493.22),
    ("IPAC_REN", "2027", "V01", 818371.34, 819979.77),
    ("ISCOM_LIL", "2027", "V01", 1844124.62, 1841951.62),
    ("ISCOM_PAR", "2027", "V01", 3048653.61, 3051354.90),
    ("ISCOM_TLS", "2027", "V01", 1709754.93, 1707811.26),
    ("MBWAY_BOR", "2027", "V01", 1863775.82, 1860582.51),
    ("MBWAY_LYO", "2027", "V01", 2620797.06, 2619135.92),
    ("MBWAY_NAN", "2027", "V01", 2267018.47, 2264136.56),
    ("MBWAY_PAR", "2027", "V01", 3292150.85, 3292759.30),
    ("PIGIER_BOR", "2027", "V01", 1191512.60, 1188191.76),
    ("PIGIER_LYO", "2027", "V01", 1653870.06, 1651005.90),
    ("TUNON_LYO", "2027", "V01", 902399.28, 905895.90),
    ("TUNON_PAR", "2027", "V01", 1133628.56, 1139411.54),
    ("IPAC_MTP", "2027", "V02", 900377.70, 907368.24),
    ("IPAC_NAN", "2027", "V02", 1174050.44, 1183473.32),
    ("IPAC_REN", "2027", "V02", 900165.39, 907368.24),
    ("ISCOM_LIL", "2027", "V02", 2043286.88, 2038255.65),
    ("ISCOM_PAR", "2027", "V02", 3373731.92, 3376549.80),
    ("ISCOM_TLS", "2027", "V02", 1893896.65, 1889819.42),
    ("MBWAY_BOR", "2027", "V02", 2069060.75, 2058872.11),
    ("MBWAY_LYO", "2027", "V02", 2907661.06, 2898267.54),
    ("MBWAY_NAN", "2027", "V02", 2515783.33, 2505434.50),
    ("MBWAY_PAR", "2027", "V02", 3651607.14, 3643681.61),
    ("PIGIER_BOR", "2027", "V02", 1317002.71, 1314822.03),
    ("PIGIER_LYO", "2027", "V02", 1826823.57, 1826960.10),
    ("TUNON_LYO", "2027", "V02", 993562.44, 1002440.79),
    ("TUNON_PAR", "2027", "V02", 1247158.51, 1260843.11),
    ("IPAC_MTP", "2027", "V03", 764876.79, 766503.19),
    ("IPAC_NAN", "2027", "V03", 997640.70, 999744.13),
    ("IPAC_REN", "2027", "V03", 764929.55, 766503.19),
    ("ISCOM_LIL", "2027", "V03", 1722442.89, 1721825.14),
    ("ISCOM_PAR", "2027", "V03", 2856528.21, 2852354.82),
    ("ISCOM_TLS", "2027", "V03", 1597221.45, 1596433.00),
    ("MBWAY_BOR", "2027", "V03", 1741134.57, 1739240.99),
    ("MBWAY_LYO", "2027", "V03", 2452416.46, 2448323.85),
    ("MBWAY_NAN", "2027", "V03", 2119379.32, 2116476.46),
    ("MBWAY_PAR", "2027", "V03", 3084114.02, 3078015.56),
    ("PIGIER_BOR", "2027", "V03", 1102717.71, 1110701.51),
    ("PIGIER_LYO", "2027", "V03", 1533137.36, 1543332.32),
    ("TUNON_LYO", "2027", "V03", 847806.31, 846816.13),
    ("TUNON_PAR", "2027", "V03", 1067029.19, 1065102.59),
]

MARQUES = ["MBWAY", "ISCOM", "IPAC", "PIGIER", "TUNON"]
CAMPUS = ["IPAC_MTP", "IPAC_NAN", "IPAC_REN", "ISCOM_LIL", "ISCOM_PAR", "ISCOM_TLS",
          "MBWAY_BOR", "MBWAY_LYO", "MBWAY_NAN", "MBWAY_PAR", "PIGIER_BOR",
          "PIGIER_LYO", "TUNON_LYO", "TUNON_PAR"]

#  La carte des lignes est calculee, pas ecrite en dur : ajouter une marque ou
#  un campus a la liste decale tout le reste sans rien casser.
L_TITRE_M = 13
L_ENT_M   = 14
L_DEB_M   = 15
L_HORS_M  = L_DEB_M + len(MARQUES)
L_TOT_M   = L_HORS_M + 1
L_TITRE_C = L_TOT_M + 3
L_ENT_C   = L_TITRE_C + 1
L_DEB_C   = L_ENT_C + 1
L_HORS_C  = L_DEB_C + len(CAMPUS)
L_TOT_C   = L_HORS_C + 1

wb = openpyxl.Workbook()

# ============================================================================
#  ONGLET 2 (cree en second, place en second) — LA ZONE DE RESTITUTION
# ============================================================================
ws = wb.active
ws.title = "Réconciliation"
dn = wb.create_sheet("Données")
dn.sheet_view.showGridLines = False
for c, w in ((1, 14), (2, 12), (3, 11), (4, 10), (5, 16), (6, 18), (7, 15)):
    dn.column_dimensions[GL(c)].width = w
#  En-tetes SANS ACCENT, et c'est la seule raison : le masque repere ses colonnes
#  par MATCH sur ce texte. Si le loader Tagetik mange un accent, le MATCH echoue
#  et tout le masque tombe en #N/A. Les libelles affiches dans les tableaux du
#  masque, eux, gardent leurs accents : ils ne servent qu'a etre lus.
ENTETES = ["Entity", "Marque", "Exercice", "Version", "CA CRM", "CA Compta", "Ecart"]
for i, t in enumerate(ENTETES):
    x = mettre(dn, 1, 1 + i, t, F(9, True), C_ if i > 3 else ind(0), fl=TETE)
    x.border = Border(top=sd(GRIS), bottom=sd(GRIS), left=sd(FIN), right=sd(FIN))
for j, (ent, an, ver, crm, cpt) in enumerate(DONNEES_LIGNES):
    r = 2 + j
    mettre(dn, r, 1, ent, F(9), ind(0))
    mettre(dn, r, 2, ent.split("_")[0], F(9), ind(0))
    mettre(dn, r, 3, an, F(9), C_)
    mettre(dn, r, 4, ver, F(9), C_)
    mettre(dn, r, 5, crm, F(9), D_, EUR)
    mettre(dn, r, 6, cpt, F(9), D_, EUR)
    mettre(dn, r, 7, "=F%d-E%d" % (r, r), F(9), D_, EURS)
    for c in range(1, 8):
        dn.cell(r, c).border = Border(bottom=sd("EDEFF2"), left=sd(FIN), right=sd(FIN))
mettre(dn, len(DONNEES_LIGNES) + 3, 1,
       "Coller ici la sortie de Q_RECONCILIATION_MARQUE. La ligne d'en-tête est "
       "repérée par son numéro, saisi en Réconciliation!C7.", F(8, False, GRIS), ind(0))
dn.freeze_panes = "A2"

# ============================================================================
#  ONGLET 1 — LE MASQUE
# ============================================================================
ws.sheet_view.showGridLines = False
NCOL = 12
for c, w in ((1, 2.5), (2, 20), (3, 16), (4, 16), (5, 15), (6, 13), (7, 11),
             (8, 2.5), (9, 22), (10, 8)):
    ws.column_dimensions[GL(c)].width = w
for r, h in {1: 6, 2: 30, 3: 6, 4: 6, 5: 20, 6: 20, 7: 20, 8: 10, 9: 15, 10: 24,
             11: 15, 12: 10, L_TITRE_M: 22, L_ENT_M: 18,
             L_TOT_M + 1: 10, L_TITRE_C: 22, L_ENT_C: 18}.items():
    ws.row_dimensions[r].height = h

# --- bandeau ---------------------------------------------------------------
for c in range(1, NCOL + 1):
    ws.cell(2, c).fill = fill(ENCRE)
mettre(ws, 2, 2, "EDUSERVICES GROUP    Réconciliation CRM / Comptabilité",
       F(12, True, "FFFFFF"), ind(0))
mettre(ws, 2, 6, "Source : Tagetik TGK_MSSQL_07", F(9, False, "C6D3DE"), D_)

# --- selecteurs ------------------------------------------------------------
def saisie(r, c, valeur, liste=None, fmt=None):
    x = mettre(ws, r, c, valeur, F(10, True), C_, fmt)
    x.fill = fill(JAUNE)
    x.border = Border(left=sd(GRIS), right=sd(GRIS), top=sd(GRIS), bottom=sd(GRIS))
    if liste:
        dv = DataValidation(type="list", formula1='"%s"' % liste, allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(x)
    return x

mettre(ws, 5, 2, "Exercice", F(9, True), ind(0))
saisie(5, 3, "2027", "2024,2025,2026,2027")
mettre(ws, 6, 2, "Version", F(9, True), ind(0))
saisie(6, 3, "V01", "ACT,V01,V02,V03")
mettre(ws, 7, 2, "Ligne d'en-tête de la restitution", F(9, True), ind(0))
saisie(7, 3, 1, None, "0")

# --- reperage des colonnes par leur nom ------------------------------------
mettre(ws, 5, 9, "Colonnes repérées par leur nom", F(8, True, GRIS), ind(0))
REPERES = [("Entity", 6), ("Marque", 7), ("Exercice", 8), ("Version", 9),
           ("CA CRM", 10), ("CA Compta", 11)]
for nom, r in REPERES:
    mettre(ws, r, 9, nom, F(8, False, GRIS), ind(0))
    mettre(ws, r, 10, '=MATCH("%s",INDEX(Données!$A:$Z,$C$7,0),0)' % nom,
           F(8, False, GRIS), C_, "0")
COL = {nom: "$J$%d" % r for nom, r in REPERES}

def colonne(nom):
    return "INDEX(Données!$A:$Z,0,%s)" % COL[nom]

def somme(mesure, cle_nom, cle_cellule):
    """Somme la mesure pour une cle, l'exercice et la version choisis."""
    return ("=SUMIFS({m},{k},{v},{ex},$C$5,{ve},$C$6)"
            .format(m=colonne(mesure), k=colonne(cle_nom), v=cle_cellule,
                    ex=colonne("Exercice"), ve=colonne("Version")))

def total(mesure):
    return ("=SUMIFS({m},{ex},$C$5,{ve},$C$6)"
            .format(m=colonne(mesure), ex=colonne("Exercice"), ve=colonne("Version")))

# --- trois chiffres cles ---------------------------------------------------
CLES = [(2, "CA CRM", "=$C$%d" % L_TOT_M), (4, "CA comptabilité", "=$D$%d" % L_TOT_M),
        (6, "Écart", "=$E$%d" % L_TOT_M)]
for c, lib, f in CLES:
    mettre(ws, 9, c, lib, F(9, True, GRIS), ind(1))
    mettre(ws, 10, c, f, F(16, True), ind(1), EURS if lib == "Écart" else EUR)
mettre(ws, 11, 6, "=IF($C${0}=0,0,$E${0}/$C${0})".format(L_TOT_M),
       F(9, False, GRIS), ind(1), PCT)
for r in range(9, 12):
    for c in range(2, 8):
        b = {}
        if r == 9: b["top"] = sd(GRIS)
        if r == 11: b["bottom"] = sd(GRIS)
        if c == 2: b["left"] = sd(GRIS)
        if c == 7: b["right"] = sd(GRIS)
        if c in (4, 6): b["left"] = sd(FIN)
        ws.cell(r, c).border = Border(**b)

# --- bloc par marque -------------------------------------------------------
def entete(r, c1, libelles, alignements):
    for i, lib in enumerate(libelles):
        x = mettre(ws, r, c1 + i, lib, F(9, True), {"g": ind(1), "d": D_, "c": C_}[alignements[i]],
                   fl=ENCRE)
        x.font = F(9, True, "FFFFFF")
        x.border = Border(top=sd(GRIS), bottom=sd(GRIS), left=sd(FIN), right=sd(FIN))

mettre(ws, L_TITRE_M, 2, "PAR MARQUE", F(10.5, True), ind(0))
for c in range(2, 8):
    ws.cell(L_TITRE_M, c).border = Border(bottom=sd(GRIS, "medium"))
entete(L_ENT_M, 2, ["Marque", "CA CRM", "CA Comptabilité", "Écart", "Écart %", ""],
       ["g", "d", "d", "d", "d", "c"])
for i, mq in enumerate(MARQUES):
    r = L_DEB_M + i
    ws.row_dimensions[r].height = 17
    mettre(ws, r, 2, mq, F(9), ind(1))
    mettre(ws, r, 3, somme("CA CRM", "Marque", "$B%d" % r), F(9), D_, EUR)
    mettre(ws, r, 4, somme("CA Compta", "Marque", "$B%d" % r), F(9), D_, EUR)
    mettre(ws, r, 5, "=D%d-C%d" % (r, r), F(9), D_, EURS)
    mettre(ws, r, 6, "=IF(C%d=0,0,E%d/C%d)" % (r, r, r), F(9), D_, PCT)
    for c in range(2, 8):
        ws.cell(r, c).border = Border(bottom=sd("EDEFF2"), left=sd(FIN), right=sd(FIN))
HORS_M = L_HORS_M
ws.row_dimensions[HORS_M].height = 17
mettre(ws, HORS_M, 2, "Hors liste", F(9, False, GRIS, True), ind(1))
mettre(ws, HORS_M, 3, "%s-SUM($C$%d:$C$%d)" % (total("CA CRM"), L_DEB_M, HORS_M - 1), F(9), D_, EUR)
mettre(ws, HORS_M, 4, "%s-SUM($D$%d:$D$%d)" % (total("CA Compta"), L_DEB_M, HORS_M - 1),
       F(9), D_, EUR)
mettre(ws, HORS_M, 5, "=D%d-C%d" % (HORS_M, HORS_M), F(9), D_, EURS)
for c in range(2, 8):
    ws.cell(HORS_M, c).border = Border(bottom=sd("EDEFF2"), left=sd(FIN), right=sd(FIN))
TOT_M = L_TOT_M
ws.row_dimensions[TOT_M].height = 19
mettre(ws, TOT_M, 2, "Total", F(9.5, True), ind(1))
for c, f in ((3, "=SUM($C$%d:$C$%d)" % (L_DEB_M, HORS_M)), (4, "=SUM($D$%d:$D$%d)" % (L_DEB_M, HORS_M)),
             (5, "=D%d-C%d" % (TOT_M, TOT_M))):
    mettre(ws, TOT_M, c, f, F(9.5, True), D_, EURS if c == 5 else EUR)
mettre(ws, TOT_M, 6, "=IF(C%d=0,0,E%d/C%d)" % (TOT_M, TOT_M, TOT_M), F(9.5, True), D_, PCT)
for c in range(2, 8):
    ws.cell(TOT_M, c).fill = fill(BLEU)
    ws.cell(TOT_M, c).border = Border(top=sd(GRIS), bottom=Side(style="double", color=NOIR))

# --- bloc par campus -------------------------------------------------------
mettre(ws, L_TITRE_C, 2, "PAR CAMPUS", F(10.5, True), ind(0))
for c in range(2, 9):
    ws.cell(L_TITRE_C, c).border = Border(bottom=sd(GRIS, "medium"))
entete(L_ENT_C, 2, ["Campus", "Marque", "CA CRM", "CA Comptabilité", "Écart", "Écart %", ""],
       ["g", "g", "d", "d", "d", "d", "c"])
for i, cp in enumerate(CAMPUS):
    r = L_DEB_C + i
    ws.row_dimensions[r].height = 17
    mettre(ws, r, 2, cp, F(9), ind(1))
    mettre(ws, r, 3, '=IF(ISERROR(FIND("_",$B%d)),$B%d,LEFT($B%d,FIND("_",$B%d)-1))'
           % (r, r, r, r), F(9, False, GRIS), ind(0))
    mettre(ws, r, 4, somme("CA CRM", "Entity", "$B%d" % r), F(9), D_, EUR)
    mettre(ws, r, 5, somme("CA Compta", "Entity", "$B%d" % r), F(9), D_, EUR)
    mettre(ws, r, 6, "=E%d-D%d" % (r, r), F(9), D_, EURS)
    mettre(ws, r, 7, "=IF(D%d=0,0,F%d/D%d)" % (r, r, r), F(9), D_, PCT)
    for c in range(2, 9):
        ws.cell(r, c).border = Border(bottom=sd("EDEFF2"), left=sd(FIN), right=sd(FIN))
HORS_C = L_HORS_C
ws.row_dimensions[HORS_C].height = 17
mettre(ws, HORS_C, 2, "Hors liste", F(9, False, GRIS, True), ind(1))
mettre(ws, HORS_C, 4, "%s-SUM($D$%d:$D$%d)" % (total("CA CRM"), L_DEB_C, HORS_C - 1), F(9), D_, EUR)
mettre(ws, HORS_C, 5, "%s-SUM($E$%d:$E$%d)" % (total("CA Compta"), L_DEB_C, HORS_C - 1),
       F(9), D_, EUR)
mettre(ws, HORS_C, 6, "=E%d-D%d" % (HORS_C, HORS_C), F(9), D_, EURS)
for c in range(2, 9):
    ws.cell(HORS_C, c).border = Border(bottom=sd("EDEFF2"), left=sd(FIN), right=sd(FIN))
TOT_C = L_TOT_C
ws.row_dimensions[TOT_C].height = 19
mettre(ws, TOT_C, 2, "Total", F(9.5, True), ind(1))
for c, f in ((4, "=SUM($D$%d:$D$%d)" % (L_DEB_C, HORS_C)), (5, "=SUM($E$%d:$E$%d)" % (L_DEB_C, HORS_C)),
             (6, "=E%d-D%d" % (TOT_C, TOT_C))):
    mettre(ws, TOT_C, c, f, F(9.5, True), D_, EURS if c == 6 else EUR)
mettre(ws, TOT_C, 7, "=IF(D%d=0,0,F%d/D%d)" % (TOT_C, TOT_C, TOT_C), F(9.5, True), D_, PCT)
for c in range(2, 9):
    ws.cell(TOT_C, c).fill = fill(BLEU)
    ws.cell(TOT_C, c).border = Border(top=sd(GRIS), bottom=Side(style="double", color=NOIR))

# --- ecarts significatifs en rouge -----------------------------------------
for zone, ref in (("E%d:E%d" % (L_DEB_M, TOT_M), "E%d" % L_DEB_M),
                  ("F%d:F%d" % (L_DEB_C, TOT_C), "F%d" % L_DEB_C)):
    ws.conditional_formatting.add(zone, FormulaRule(
        formula=["ABS(%s)>1" % ref], font=Font(name=UI, size=9, bold=True, color=ROUGE)))

mettre(ws, TOT_C + 3, 2,
       "CA CRM = effectifs x tarif + nouveaux x frais d'inscription (socle, ou moteur "
       "pour 2027). CA Comptabilité = comptes 706 + 7062 + 708.",
       F(8, False, GRIS), ind(0))
ws.freeze_panes = "B%d" % L_DEB_M
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize = 9
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_area = "A1:G%d" % (TOT_C + 3)

# ============================================================================
#  NORMALISATION ET CONTROLES
# ============================================================================
for feuille, nl, nc in ((ws, TOT_C + 4, NCOL), (dn, len(DONNEES_LIGNES) + 4, 7)):
    for r in range(1, nl + 1):
        for c in range(1, nc + 1):
            x = feuille.cell(r, c)
            if x.font is None or x.font.name != UI:
                x.font = F(9)
            if not x.number_format:
                x.number_format = "General"

wb.save(OUT)
v = openpyxl.load_workbook(OUT)
pb = []
for f in v.worksheets:
    if f.merged_cells.ranges:
        pb.append("%s : %d fusion(s)" % (f.title, len(f.merged_cells.ranges)))
    for row in f.iter_rows():
        for cel in row:
            if cel.number_format == "":
                pb.append("%s!%s : format vide" % (f.title, cel.coordinate))
            if isinstance(cel.value, str) and "_xlfn" in cel.value:
                pb.append("%s!%s : _xlfn" % (f.title, cel.coordinate))
print("onglets  :", v.sheetnames)
print("lignes de donnees :", len(DONNEES_LIGNES))
print("fusions  :", sum(len(f.merged_cells.ranges) for f in v.worksheets))
print("controle :", "PASS" if not pb else pb[:5])

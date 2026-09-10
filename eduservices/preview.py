#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Rendu fidele (matplotlib) des onglets redesignes en lisant les fills/valeurs."""
import openpyxl, matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch
from openpyxl.utils import column_index_from_string as ci

wb = openpyxl.load_workbook("CAD_SAAD_LIVE.xlsx")

def argb(cell):
    f = cell.fill
    if f and f.fgColor and f.fgColor.rgb and isinstance(f.fgColor.rgb, str):
        rgb = f.fgColor.rgb
        if len(rgb) == 8: rgb = rgb[2:]
        if rgb not in ("000000",) or f.patternType == "solid":
            if f.patternType == "solid": return "#" + rgb
    return None

def fontcol(cell):
    c = cell.font.color
    if c and c.rgb and isinstance(c.rgb, str):
        rgb = c.rgb[2:] if len(c.rgb) == 8 else c.rgb
        return "#" + rgb
    return "#000000"

def fmt_val(cell):
    v = cell.value
    if v is None or (isinstance(v, str) and v.startswith("=")): return ""
    if isinstance(v, (int, float)):
        nf = cell.number_format
        if "%" in nf: return "%.1f%%" % (v * 100)
        if "0.00" in nf: return "%.2f" % v
        if "##0" in nf: return "{:,.0f}".format(v).replace(",", " ")
        return str(v)
    return str(v)

def render(sheet, r0, r1, c0, c1, colw, title, out, charts):
    ws = wb[sheet]
    ncol = c1 - c0 + 1
    xs = [0]
    for c in range(c0, c1 + 1):
        xs.append(xs[-1] + colw.get(c, 1.4))
    W = xs[-1]
    H = (r1 - r0 + 1) * 0.34
    fig, ax = plt.subplots(figsize=(W * 0.9, H * 0.9 + 1.2))
    ax.set_xlim(0, W); ax.set_ylim(0, H + 0.9); ax.axis("off")
    ax.add_patch(Rectangle((0, H + 0.35), W, 0.55, color="#1F3864"))
    ax.text(0.15, H + 0.62, title, color="white", fontsize=13, fontweight="bold", va="center")
    for ri, r in enumerate(range(r0, r1 + 1)):
        y = H - (ri + 1) * 0.34
        for ci_, c in enumerate(range(c0, c1 + 1)):
            cell = ws.cell(r, c)
            x = xs[ci_]; w = xs[ci_ + 1] - xs[ci_]
            bg = argb(cell)
            if bg:
                ax.add_patch(Rectangle((x, y), w, 0.34, facecolor=bg, edgecolor="#D9D9D9", lw=0.4))
            else:
                ax.add_patch(Rectangle((x, y), w, 0.34, facecolor="none", edgecolor="#EEEEEE", lw=0.2))
            t = fmt_val(cell)
            if t:
                fc = fontcol(cell)
                bold = cell.font.bold
                ha = "left" if (cell.alignment.horizontal in (None, "left")) else \
                     ("right" if cell.alignment.horizontal == "right" else "center")
                tx = x + 0.06 if ha == "left" else (x + w - 0.06 if ha == "right" else x + w / 2)
                ax.text(tx, y + 0.17, t[:34], color=fc, fontsize=7.5,
                        fontweight="bold" if bold else "normal", ha=ha, va="center")
    # placeholders graphes
    for (cx0, cx1, ry0, ry1, label) in charts:
        x = xs[cx0 - c0]; xe = xs[min(cx1, c1) - c0 + 1] if cx1 <= c1 else W
        ytop = H - (ry0 - r0) * 0.34; ybot = H - (ry1 - r0 + 1) * 0.34
        ax.add_patch(FancyBboxPatch((x + 0.05, ybot), xe - x - 0.1, ytop - ybot,
                     boxstyle="round,pad=0.02", facecolor="#F4F8FC", edgecolor="#2E75B6", lw=1.2))
        ax.text((x + xe) / 2, (ytop + ybot) / 2, "📊 " + label, color="#2E75B6",
                fontsize=9, ha="center", va="center", fontweight="bold")
    plt.tight_layout()
    plt.savefig(out, dpi=115, bbox_inches="tight")
    print("wrote", out)

# --- cad ---
colw_cad = {ci("G"):3.2, ci("H"):5.2, ci("I"):1.7, ci("J"):1.7, ci("K"):1.7,
            ci("L"):0.3, ci("M"):1.5, ci("N"):1.5, ci("O"):1.5, ci("P"):1.5, ci("Q"):1.5, ci("R"):1.5}
render("cad", 2, 57, ci("G"), ci("R"), colw_cad,
       "CADRAGE  ·  Hypotheses budgetaires 2027",
       "/tmp/preview_cad.png",
       [(ci("M"), ci("R"), 8, 18, "Coef prix / marque"),
        (ci("M"), ci("R"), 20, 34, "Parametres par version (V01/V02/V03)")])

# --- Pilotage ---
colw_pil = {ci("D"):2.6, ci("E"):2.6, ci("F"):2.1, ci("G"):1.5, ci("H"):1.4, ci("I"):1.4,
            ci("J"):1.3, ci("K"):1.3, ci("L"):1.3, ci("M"):1.5, ci("N"):1.6, ci("O"):1.8,
            ci("P"):0.3, ci("Q"):1.7, ci("R"):1.4, ci("S"):1.4, ci("T"):1.4, ci("U"):1.4}
render("Pilotage", 2, 42, ci("D"), ci("U"), colw_pil,
       "PILOTAGE  ·  Cap strategique & Cles d'allocation",
       "/tmp/preview_pil.png",
       [(ci("D"), ci("L"), 28, 42, "Budget reference vs rejoue (live)"),
        (ci("N"), ci("U"), 28, 42, "Cap retenu / campus")])

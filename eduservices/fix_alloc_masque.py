# -*- coding: utf-8 -*-
"""Correctif du masque d'allocation, appliqué SUR LE FICHIER LUI-MÊME.

Patch XML, jamais un aller-retour openpyxl : ce classeur porte _TGK_HIDDEN
(veryHidden), xl/webextensions, customXml, customProperty1.bin et quatre plages
nommées ALLOC_* que Tagetik lit pour ses saisies. Un aller-retour en perdrait
une partie sans le dire.
"""
import re, zipfile, shutil, collections, warnings
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.formula.translate import Translator

SRC, DST = "src.xlsm", "ALLOC_corrige.xlsm"
DEB, GRP, FIN_OLD, FIN_NEW = 18, 97, 349, 420
S_ALLOC, S_CALC = "xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml"

ff = openpyxl.load_workbook("clean.xlsx")
a, c = ff["Allocation"], ff["_CALC_ALLOC"]

MARQUE = {"MBway": "MBWAY", "ISCOM": "ISCOM", "Ipac Bachelor Factory": "IPAC",
          "Pigier": "PIGIER", "Tunon": "TUNON"}
VILLE = {"Paris": "PAR", "Lyon": "LYO", "Nantes": "NAN", "Bordeaux": "BOR",
         "Lille": "LIL", "Toulouse": "TLS", "Rennes": "REN", "Montpellier": "MTP"}
PROG = {"Bachelor Commerce": "BAC_CCE", "Bachelor Communication": "BAC_COM",
        "Bachelor Management": "BAC_MGT", "Bachelor Ressources humaines": "BAC_RH",
        "Bachelor Tourisme": "BAC_TOU", "BTS Gestion": "BTS_GES",
        "Mastère Communication": "MAS_COM", "Mastère Management": "MAS_MGT"}
ANS = {"BAC": {"1re année": "B1", "2e année": "B2", "3e année": "B3"},
       "MAS": {"1re année": "M1", "2e année": "M2"},
       "BTS": {"1re année": "BTS1", "2e année": "BTS2"}}
MOD = {"initial": "INIT", "alternance": "ALT"}

#  Le niveau se lit dans la FORMULE, restée juste, et non dans le plan de la
#  feuille : celui-ci a dérivé sur dix lignes à partir de la 81.
def niveau(r):
    f = a.cell(r, 3).value or ""
    if "$V%d" % r in f: return 1
    if "$W%d" % r in f: return 3
    if "$U%d" % r in f: return 2
    return 0

CLES, mq, cp, classes, campus = {}, None, None, [], []
for r in range(DEB, GRP):
    lib = (a.cell(r, 2).value or "").strip()
    niv = niveau(r)
    if niv == 1:
        mq = MARQUE[lib]; cp = None
        vals = (mq, "", mq, "", "", "")
    elif niv == 2:
        cp = "%s_%s" % (mq, VILLE[lib]); campus.append(cp)
        vals = (cp, cp, mq, "", "", "")
    else:
        pl, al, ml = [t.strip() for t in lib.split(",")]
        p = PROG[pl]; an = ANS[p.split("_")[0]][al]; mo = MOD[ml]
        classes.append((cp, p, an, mo))
        vals = ("|".join((cp, p, an, mo)), cp, mq, p, an, mo)
    CLES[r] = (vals, niv)
CLES[GRP] = (("GROUPE", "", "", "", "", ""), 0)

n = collections.Counter(niv for _, niv in CLES.values())
assert (n[1], n[2], n[3]) == (5, 14, 60), (n[1], n[2], n[3])
assert len(set(classes)) == 60 and len(set(campus)) == 14
print("plan : %d marques, %d campus, %d classes, toutes distinctes" % (n[1], n[2], n[3]))

MODELE = [(c.cell(FIN_OLD, col).column_letter, c.cell(FIN_OLD, col).value)
          for col in range(1, 47) if c.cell(FIN_OLD, col).value is not None]

def esc(t):
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

zin = zipfile.ZipFile(SRC); tmp = DST + ".tmp"
zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)

    if it.filename == S_ALLOC:
        s = data.decode("utf-8")
        if 'min="20" max="26"' not in s:
            s = s.replace("<cols>", '<cols><col min="20" max="26" width="13" '
                                    'hidden="1" customWidth="1"/>', 1)
        for r, (vals, niv) in CLES.items():
            m = re.search(r'<row r="%d"[^>]*>' % r, s)
            bal = re.sub(r'\s+(?:outlineLevel="\d+"|hidden="\d")', "", m.group(0))
            bal = bal.replace('spans="2:13"', 'spans="2:26"')
            sup = (' outlineLevel="%d"' % (niv - 1)) if niv > 1 else ""
            if niv == 3:
                sup += ' hidden="1"'
            s = s[:m.start()] + bal[:-1] + sup + ">" + s[m.end():]
            cells = "".join('<c r="%s%d" t="inlineStr"><is><t>%s</t></is></c>'
                            % (L, r, esc(x)) for L, x in zip("TUVWXY", vals) if x != "")
            cells += '<c r="Z%d"><v>%d</v></c>' % (r, niv)
            s = re.sub(r'(<c r="M%d"(?:[^>]*?/>|[^>]*?>.*?</c>))' % r,
                       lambda mm: mm.group(1) + cells, s, count=1, flags=re.S)

        #  C12 portait 1000 : un nombre, qu'aucune branche de P1 ne reconnaît.
        #  P1 retombait donc silencieusement sur V01, quelle que soit la liste.
        s = re.sub(r'<c r="C12"[^>]*(?:/>|>.*?</c>)',
                   '<c r="C12" s="5" t="inlineStr"><is><t>Cadrage</t></is></c>'
                   '<c r="D12"><f>"version "&amp;$P$1&amp;"  ·  "&amp;TEXT(COUNTIF('
                   '_CALC_ALLOC!$AS$1:$AS$%d,"2027|"&amp;$P$1),"0")&amp;'
                   '" classes lues sur 60"</f></c>' % FIN_NEW, s, count=1, flags=re.S)

        #  P1 lit désormais les mêmes cellules que la liste déroulante : le menu
        #  et le décodage ne peuvent plus diverger, et une valeur inattendue
        #  rend « ? » -- donc un tableau à zéro, visible.
        s = s.replace('IF($C$12="Optimiste","V02",IF($C$12="Prudent","V03","V01"))',
                      'IF($C$12=$E$5,"V01",IF($C$12=$E$6,"V02",IF($C$12=$E$7,"V03","?")))')
        data = s.encode("utf-8")

    elif it.filename == S_CALC:
        s = data.decode("utf-8")
        blocs = []
        for r in range(FIN_OLD + 1, FIN_NEW + 1):
            cs = "".join('<c r="%s%d"><f>%s</f></c>'
                         % (L, r, esc(Translator(f, origin="%s%d" % (L, FIN_OLD))
                                      .translate_formula("%s%d" % (L, r))[1:]))
                         for L, f in MODELE if isinstance(f, str) and f.startswith("="))
            blocs.append('<row r="%d" spans="1:45">%s</row>' % (r, cs))
        s = s.replace("</sheetData>", "".join(blocs) + "</sheetData>")
        s = s.replace('<dimension ref="A1:AT%d"/>' % FIN_OLD,
                      '<dimension ref="A1:AT%d"/>' % FIN_NEW)
        data = s.encode("utf-8")

    elif it.filename == "xl/workbook.xml":
        s = data.decode("utf-8")
        s = re.sub(r'<calcPr[^/]*/>', '<calcPr calcId="0" fullCalcOnLoad="1"/>', s)
        data = s.encode("utf-8")

    zout.writestr(it, data)
zout.close(); zin.close()
shutil.move(tmp, DST)
print("ecrit :", DST)

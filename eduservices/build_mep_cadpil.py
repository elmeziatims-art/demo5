#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_mep_cadpil.py — mise en page de CAD_PIL (Cadrage + Pilotage).

Jumeau visuel des deux onglets, a geometrie constante : memes lignes, memes
colonnes que le fichier reel, pour que le collage special "Formats" tombe juste.
Les chiffres sont figes (valeurs en cache du fichier reel).

PARTI PRIS : un fichier de controle de gestion, pas une maquette. Fond blanc,
filets fins, un seul bleu, saisies en jaune pale, negatifs en rouge, double
trait sous les totaux. Aucun commentaire dans les cellules.
"""
import warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (RichTextProperties, Paragraph, ParagraphProperties,
                                   CharacterProperties, Font as PoliceDessin)
from openpyxl.utils import get_column_letter as GL
warnings.filterwarnings("ignore")

OUT = "/home/user/demo5/eduservices/CAD_PIL_MEP.xlsx"

NOIR   = "000000"
ENCRE  = "1F3B57"     # bandeau de titre
TETE   = "D9E2EC"     # en-tetes de tableau
BLEU   = "DDEBF7"     # colonne retenue, lignes de synthese
JAUNE  = "FFF2CC"     # cellules de saisie
GRIS   = "595959"     # filets structurants
FIN    = "BFBFBF"     # filets de tableau
CLAIR  = "F2F2F2"     # sous-totaux
ROUGE  = "C00000"
UI = "Arial"

def F(sz=9, b=False, c=NOIR, i=False):
    return Font(name=UI, size=sz, bold=b, color=c, italic=i)
def fill(c):
    return PatternFill("solid", fgColor=c)
def sd(c=FIN, st="thin"):
    return Side(style=st, color=c)

G  = Alignment("left",   vertical="center")
Dr = Alignment("right",  vertical="center")
Ce = Alignment("center", vertical="center")
def ind(n=1):
    return Alignment("left", vertical="center", indent=n)

EUR  = '#,##0" €";[Red]-#,##0" €"'
PCT  = '0.0%;[Red]-0.0%'
TAUX = '0.00%;[Red]-0.00%'
NB   = '#,##0;[Red]-#,##0'
NBD  = '#,##0.0'
COEF = '0.00'

def peindre(ws, r1, c1, r2, c2, couleur):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill(couleur)

def mettre(ws, r, c, v, font=None, al=None, fmt=None, fl=None, bd=None):
    x = ws.cell(r, c, v)
    x.font = font or F()
    x.alignment = al or G
    x.number_format = fmt or "General"
    if fl: x.fill = fill(fl)
    if bd: x.border = bd
    return x

def bandeau(ws, ncol, titre, mention, cmention):
    peindre(ws, 2, 1, 2, ncol, ENCRE)
    mettre(ws, 2, 2, titre, F(12, True, "FFFFFF"), ind(0))
    mettre(ws, 2, cmention, mention, F(9, False, "C6D3DE"), Dr)
    ws.row_dimensions[1].height = 6.0
    ws.row_dimensions[2].height = 30.0
    ws.row_dimensions[3].height = 6.0

def titre_section(ws, r, c1, c2, texte):
    mettre(ws, r, c1, texte, F(10.5, True, NOIR), ind(0))
    for c in range(c1, c2 + 1):
        ws.cell(r, c).border = Border(bottom=sd(GRIS, "medium"))

def entete(ws, r, c1, libelles, alignements):
    for i, lib in enumerate(libelles):
        x = mettre(ws, r, c1 + i, lib, F(9, True, NOIR),
                   {"g": ind(1), "d": Dr, "c": Ce}[alignements[i]], fl=TETE)
        x.border = Border(top=sd(GRIS), bottom=sd(GRIS),
                          left=sd(FIN), right=sd(FIN))

def encadrer(ws, r1, c1, r2, c2, coupes=()):
    """Un cadre fin, avec des refends verticaux aux colonnes indiquees."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            b = {}
            if r == r1: b["top"] = sd(GRIS)
            if r == r2: b["bottom"] = sd(GRIS)
            if c == c1: b["left"] = sd(GRIS)
            if c == c2: b["right"] = sd(GRIS)
            if c in coupes: b["left"] = sd(FIN)
            ws.cell(r, c).border = Border(**b)

def ligne_total(ws, r, c1, c2, fond=CLAIR, double=False):
    for c in range(c1, c2 + 1):
        ws.cell(r, c).fill = fill(fond)
        ws.cell(r, c).border = Border(top=sd(GRIS),
                                      bottom=Side(style="double" if double else "thin",
                                                  color=NOIR if double else FIN))

def saisie(ws, r, c):
    x = ws.cell(r, c)
    x.fill = fill(JAUNE)
    x.border = Border(left=sd(GRIS), right=sd(GRIS), top=sd(GRIS), bottom=sd(GRIS))
    return x

def impression(ws, lignes_titre, zone):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = lignes_titre
    ws.print_area = zone
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5

wb = openpyxl.Workbook()

# ============================================================================
#  CADRAGE   —   carte des lignes identique au fichier reel
#    2 titre · 5 scenario · 7-9 chiffres cles · 10 section · 11 entete
#    12-15 reconciliation · 18 section · 19 entete · 21 sous-titre croissance
#    22 ligne technique du MATCH · 23-28 leviers · 32 sous-titre couts
#    34-38 leviers · 40 sous-titre constante · 42 frais de dossier
#    J21:K28 coefficients de prix
# ============================================================================
ws = wb.active
ws.title = "Cadrage"
ws.sheet_view.showGridLines = False
NCOL = 19
for c, w in ((1, 2.0), (2, 46.0), (3, 14.5), (4, 14.5), (5, 14.5), (6, 14.5),
             (7, 2.0), (8, 2.0), (9, 2.0), (10, 25.0), (11, 9.5), (12, 2.0), (13, 12.0),
             (16, 8.0), (17, 14.0), (18, 12.0), (19, 12.0), (20, 11.0)):
    ws.column_dimensions[GL(c)].width = w
for r, h in {4: 6.0, 5: 20.0, 6: 12.0, 7: 15.0, 8: 24.0, 9: 15.0, 10: 24.0, 11: 18.0,
             12: 16.5, 13: 16.5, 14: 16.5, 15: 16.5, 16: 12.0, 17: 6.0, 18: 24.0,
             19: 18.0, 20: 6.0, 21: 18.0, 22: 9.0, 23: 15.75, 24: 15.75, 25: 15.75,
             26: 15.75, 27: 15.75, 28: 15.75, 29: 6.0, 30: 6.0, 31: 6.0, 32: 18.0,
             33: 6.0, 34: 15.75, 35: 15.75, 36: 15.75, 37: 15.75, 38: 15.75, 39: 6.0,
             40: 18.0, 41: 6.0, 42: 15.75, 43: 12.0, 44: 12.0}.items():
    ws.row_dimensions[r].height = h

bandeau(ws, NCOL, "EDUSERVICES GROUP    Cadrage budgetaire 2027".replace("budgetaire", "budgétaire"),
        "Scénario 2027BUD_V1", 11)

mettre(ws, 5, 2, "Scénario retenu", F(9, True), ind(0))
mettre(ws, 5, 3, "Optimiste", F(10, True), Ce)
saisie(ws, 5, 3)
mettre(ws, 5, 5, "Objectif de croissance EBITDA", F(9, True), Dr)
mettre(ws, 5, 6, 0.065, F(10, True), Ce, PCT)
saisie(ws, 5, 6)
dv = DataValidation(type="list", formula1='"Cadrage,Optimiste,Prudent"', allow_blank=False)
ws.add_data_validation(dv); dv.add(ws["C5"])

# --- chiffres cles ---------------------------------------------------------
CLES = [(2, 2, "Objectif EBITDA 2027", 4095766.36, EUR, "2026 : 3 845 790 €"),
        (3, 4, "EBITDA construit", 6744302.91, EUR, "marge 25,2 %"),
        (5, 6, "Écart à l'objectif", 2648536.55, EUR, "objectif dépassé")]
peindre(ws, 7, 2, 9, 6, "FFFFFF")
for c1, c2, lib, val, fmt, sous in CLES:
    mettre(ws, 7, c1, lib, F(9, True, GRIS), ind(1))
    mettre(ws, 8, c1, val, F(16, True, NOIR), ind(1), fmt)
    mettre(ws, 9, c1, sous, F(8, False, GRIS), ind(1))
encadrer(ws, 7, 2, 9, 6, coupes=(3, 5))

# --- reconciliation --------------------------------------------------------
titre_section(ws, 10, 2, 6, "RÉCONCILIATION 2026 / 2027")
entete(ws, 11, 2, ["Indicateur", "Réel 2026", "Budget 2027", "Écart", "Écart %"],
       ["g", "d", "d", "d", "d"])
REC = [("Chiffre d'affaires", 23098985, 26814168.68, 3715183.68, 0.1608, EUR),
       ("EBITDA", 3845790.01, 6744302.91, 2898512.90, 0.7537, EUR),
       ("Marge EBITDA", 0.1665, 0.2515, 0.0850, None, PCT),
       ("Effectif moyen", 3114, 3473.67, 359.67, 0.1155, NB)]
for i, (lib, a, b, e, pc, f1) in enumerate(REC):
    r = 12 + i
    mettre(ws, r, 2, lib, F(9, i == 1), ind(1))
    mettre(ws, r, 3, a, F(9, i == 1), Dr, f1)
    mettre(ws, r, 4, b, F(9, i == 1), Dr, f1)
    mettre(ws, r, 5, e, F(9, i == 1), Dr, f1)
    if pc is not None:
        mettre(ws, r, 6, pc, F(9, i == 1), Dr, PCT)
    for c in range(2, 7):
        ws.cell(r, c).border = Border(bottom=sd(FIN), left=sd(FIN), right=sd(FIN))
ligne_total(ws, 13, 2, 6, BLEU)

# --- leviers ---------------------------------------------------------------
titre_section(ws, 18, 2, 6, "HYPOTHÈSES DE CONSTRUCTION")
entete(ws, 19, 2, ["Levier", "Cadrage", "Optimiste", "Prudent", "Retenu"],
       ["g", "c", "c", "c", "c"])
ws.cell(19, 6).fill = fill(BLEU)

def sous_titre(r, texte):
    mettre(ws, r, 2, texte, F(9, True, ENCRE), ind(0))
    for c in range(2, 7):
        ws.cell(r, c).border = Border(bottom=sd(FIN))

def leviers(r0, lignes, fmt):
    for i, (lib, cadr, opti, prud, ret) in enumerate(lignes):
        r = r0 + i
        mettre(ws, r, 2, lib, F(9), ind(1))
        for j, v in enumerate((cadr, opti, prud)):
            if v is None:
                mettre(ws, r, 3 + j, "n.d.", F(9, False, ROUGE), Ce)
            else:
                mettre(ws, r, 3 + j, v, F(9), Ce, fmt)
        mettre(ws, r, 6, ret, F(9, True), Ce, fmt, fl=BLEU)
        for c in range(2, 7):
            ws.cell(r, c).border = Border(bottom=sd(FIN), left=sd(FIN), right=sd(FIN))

sous_titre(21, "Croissance")
for j, lib in enumerate(("Cadrage", "Optimiste", "Prudent")):   # ligne technique du MATCH
    mettre(ws, 22, 3 + j, lib, F(6, False, "FFFFFF"), Ce)
leviers(23, [("Variation du budget d'acquisition", .08, .15, -.05, .15),
             ("Variation du budget de marque", .10, .15, -.05, .15),
             ("Hausse tarifaire", .0029, .035, .02, .035),
             ("Gain de conversion lead / candidature", .01, .03, -.01, .03),
             ("Gain de conversion admis / inscrit", .01, .025, -.01, .025),
             ("Amélioration du taux de passage", .005, .015, -.01, .015)], TAUX)
sous_titre(32, "Coûts")
leviers(34, [("Inflation des charges externes", .02, .015, .03, .015),
             ("Politique salariale", .025, .02, .03, .02),
             ("Variation des effectifs permanents", .04, .03, .05, .03),
             ("Effort de productivité", .0185, .03, None, .03),
             ("Variation des coûts de structure", None, -.03, .04, -.03)], TAUX)
sous_titre(40, "Constante")
leviers(42, [("Frais de dossier par nouvel inscrit", 90, 90, 90, 90)], EUR)

mettre(ws, 44, 2, "Source : Tagetik TGK_MSSQL_07, cube d'hypothèses, versions V01 / V02 / V03.",
       F(8, False, GRIS), ind(0))

# --- coefficients de prix --------------------------------------------------
titre_section(ws, 21, 10, 11, "COEFFICIENTS DE PRIX")
entete(ws, 23, 10, ["Marque", "Coeff."], ["g", "c"])
for i, (m, k) in enumerate([("MBway", 1.20), ("ISCOM", 1.15),
                            ("Ipac Bachelor Factory", 0.95), ("Pigier", 0.90),
                            ("Tunon", 1.05)]):
    r = 24 + i
    mettre(ws, r, 10, m, F(9), ind(1))
    mettre(ws, r, 11, k, F(9), Ce, COEF)
    for c in (10, 11):
        ws.cell(r, c).border = Border(bottom=sd(FIN), left=sd(FIN), right=sd(FIN))

# --- zone de donnees des graphiques (hors zone d'impression) ---------------
#  Reel 2024-2026 reconstitue depuis l'onglet PNL : comptes 70x pour le chiffre
#  d'affaires, moins les 6xx hors 6811 pour l'EBITDA. Controle : 2026 retombe
#  sur 23 098 985 EUR et 3 845 790 EUR.
mettre(ws, 4, 16, "Données des graphiques", F(8, True, GRIS), ind(0))
for i, lib in enumerate(("Exercice", "Chiffre d'affaires", "EBITDA",
                         "Objectif EBITDA", "Marge EBITDA")):
    mettre(ws, 5, 16 + i, lib, F(8, True, GRIS), Ce if i else ind(0))
GRAPHE = [(2024, 20567210, 3151035, 4095766.36, .1532),
          (2025, 21758770, 3467768, 4095766.36, .1594),
          (2026, 23098985, 3845790, 4095766.36, .1665),
          (2027, 26814168.68, 6744302.91, 4095766.36, .2515)]
for i, (an, ca, eb, obj, mg) in enumerate(GRAPHE):
    r = 6 + i
    mettre(ws, r, 16, an, F(8, False, GRIS), Ce, "0")
    for j, (v, f) in enumerate(((ca, NB), (eb, NB), (obj, NB), (mg, PCT))):
        mettre(ws, r, 17 + j, v, F(8, False, GRIS), Dr, f)

police = CharacterProperties(latin=PoliceDessin(typeface=UI), sz=800, solidFill=NOIR)
def texte_arial(taille=800):
    cp = CharacterProperties(latin=PoliceDessin(typeface=UI), sz=taille, solidFill=NOIR)
    return RichText(bodyPr=RichTextProperties(),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])
def titrer(graphe, texte):
    gras = CharacterProperties(latin=PoliceDessin(typeface=UI), sz=1000, b=True,
                               solidFill=NOIR)
    graphe.title = texte
    graphe.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=gras)
    graphe.title.tx.rich.p[0].r[0].rPr = gras

def point_budget(serie, couleur):
    """La colonne 2027 n'est pas un realise : elle se distingue."""
    pt = DataPoint(idx=3)
    pt.graphicalProperties.solidFill = couleur
    serie.data_points = [pt]

#  --- graphique 1 : le chiffre d'affaires et l'EBITDA, meme axe ------------
volumes = BarChart()
volumes.type = "col"
volumes.grouping = "clustered"
volumes.gapWidth = 60
volumes.overlap = -15
volumes.add_data(Reference(ws, min_col=17, max_col=18, min_row=5, max_row=9),
                 titles_from_data=True)
volumes.set_categories(Reference(ws, min_col=16, min_row=6, max_row=9))
volumes.series[0].graphicalProperties.solidFill = "9DC3E6"      # chiffre d'affaires
volumes.series[1].graphicalProperties.solidFill = ENCRE         # EBITDA
point_budget(volumes.series[0], "C9DEF2")
point_budget(volumes.series[1], "2E75B6")
volumes.dLbls = DataLabelList()
volumes.dLbls.showVal = True
#  huit etiquettes sur 10 cm : en euros pleins elles se chevauchent, en millions
#  a une decimale elles tiennent et restent lisibles.
volumes.dLbls.numFmt = '#,##0.0,," M€"'
volumes.dLbls.txPr = texte_arial(700)

objectif = LineChart()
objectif.add_data(Reference(ws, min_col=19, min_row=5, max_row=9), titles_from_data=True)
objectif.series[0].graphicalProperties.line.solidFill = ROUGE
objectif.series[0].graphicalProperties.line.dashStyle = "dash"
objectif.series[0].graphicalProperties.line.width = 18000
objectif.series[0].smooth = False
objectif.y_axis.axId = volumes.y_axis.axId
objectif.x_axis.axId = volumes.x_axis.axId
volumes += objectif

titrer(volumes, "Chiffre d'affaires et EBITDA — réel 2024-2026, budget 2027")
volumes.y_axis.numFmt = '#,##0" €"'
volumes.x_axis.numFmt = "0"
volumes.x_axis.delete = False
volumes.y_axis.delete = False
volumes.x_axis.axPos = "b"
volumes.y_axis.axPos = "l"
volumes.legend.position = "b"
volumes.legend.overlay = False
volumes.txPr = texte_arial()
volumes.height = 7.4
volumes.width = 10.4
ws.add_chart(volumes, "H5")

#  --- graphique 2 : la marge, que l'axe en euros ne peut pas montrer -------
marge = BarChart()
marge.type = "col"
marge.grouping = "clustered"
marge.gapWidth = 90
marge.add_data(Reference(ws, min_col=20, min_row=5, max_row=9), titles_from_data=True)
marge.set_categories(Reference(ws, min_col=16, min_row=6, max_row=9))
marge.series[0].graphicalProperties.solidFill = ENCRE
point_budget(marge.series[0], "2E75B6")
marge.dLbls = DataLabelList()
marge.dLbls.showVal = True
marge.dLbls.numFmt = '0.0%'
marge.dLbls.txPr = texte_arial()
titrer(marge, "Marge EBITDA")
marge.y_axis.numFmt = "0%"
marge.x_axis.numFmt = "0"
marge.x_axis.delete = False
marge.y_axis.delete = False
marge.x_axis.axPos = "b"
marge.y_axis.axPos = "l"
marge.legend = None
marge.txPr = texte_arial()
marge.height = 4.6
marge.width = 10.4
ws.add_chart(marge, "H31")

ws.freeze_panes = "B10"
impression(ws, "1:3", "A1:M44")

# ============================================================================
#  PILOTAGE
# ============================================================================
p = wb.create_sheet("Pilotage")
p.sheet_view.showGridLines = False
NC2 = 23
for c, w in ((1, 3.0), (2, 16.5), (3, 13.5), (4, 12.5), (5, 13.0), (6, 12.0),
             (7, 10.5), (8, 13.0), (9, 10.5), (10, 12.0), (11, 13.0), (12, 12.5),
             (13, 2.0)):
    p.column_dimensions[GL(c)].width = w
for lettre in ("N", "U", "V", "W"):
    p.column_dimensions[lettre].hidden = True
for r, h in {4: 6.0, 5: 6.0, 6: 20.0, 7: 6.0, 8: 6.0, 9: 15.0, 10: 24.0, 11: 15.0,
             12: 12.0, 13: 6.0, 14: 6.0, 15: 6.0, 16: 24.0, 17: 6.0, 18: 26.0,
             33: 12.0, 34: 6.0, 35: 6.0, 36: 24.0, 37: 6.0, 38: 26.0,
             53: 17.0, 54: 15.75, 55: 19.0, 56: 12.0, 57: 12.0}.items():
    p.row_dimensions[r].height = h
for r in list(range(19, 33)) + list(range(39, 53)):
    p.row_dimensions[r].height = 15.75

bandeau(p, NC2, "EDUSERVICES GROUP    Pilotage budgétaire 2027", "Scénario 2027BUD_V1", 12)

mettre(p, 6, 2, "Scénario retenu", F(9, True), ind(0))
mettre(p, 6, 5, "Cadrage", F(10, True), Ce)
saisie(p, 6, 5)
dv2 = DataValidation(type="list", formula1='"Cadrage,Optimiste,Prudent"', allow_blank=False)
p.add_data_validation(dv2); dv2.add(p["E6"])

CLES2 = [(2, 2, "Chiffre d'affaires", 24231703.55, EUR, "2026 : 23 098 985 €"),
         (3, 4, "EBITDA après siège", 4111503.64, EUR, "2026 : 3 845 790 €"),
         (5, 6, "Marge EBITDA", 0.1697, PCT, "2026 : 16,6 %"),
         (7, 8, "Effectif moyen", 3253.62, NBD, "2026 : 3 114"),
         (9, 11, "Croissance CA", 0.0490, PCT, "vs réel 2026")]
peindre(p, 9, 2, 11, 11, "FFFFFF")
for c1, c2, lib, val, fmt, sous in CLES2:
    mettre(p, 9, c1, lib, F(9, True, GRIS), ind(1))
    mettre(p, 10, c1, val, F(16, True, NOIR), ind(1), fmt)
    mettre(p, 11, c1, sous, F(8, False, GRIS), ind(1))
encadrer(p, 9, 2, 11, 11, coupes=(3, 5, 7, 9))

CAP = [("IPAC_MTP", "Ipac", "Montpellier", 1033.98, .1448, .0194, 1.1626, 1.3855, .9443, 1, 15178),
       ("IPAC_NAN", "Ipac", "Nantes", 1216.29, .1379, .0213, .9883, 1.3195, .8591, 1, 21760),
       ("IPAC_REN", "Ipac", "Rennes", 1040.56, .1448, .0191, 1.1553, 1.3855, .9598, 1, 14933),
       ("ISCOM_LIL", "ISCOM", "Lille", 978.58, .0815, .0144, 1.2284, .7804, 1.2767, 1, 25220),
       ("ISCOM_PAR", "ISCOM", "Paris", 1540.33, .0777, .0209, .7804, .7440, .8773, 1, 60800),
       ("ISCOM_TLS", "ISCOM", "Toulouse", 919.03, .0821, .0136, 1.3080, .7858, 1.3518, 1, 22083),
       ("MBWAY_BOR", "MBway", "Bordeaux", 889.59, .1261, .0141, 1.3513, 1.2067, 1.3049, 1, 24923),
       ("MBWAY_LYO", "MBway", "Lyon", 1101.60, .1189, .0167, 1.0912, 1.1376, 1.0957, 1, 41783),
       ("MBWAY_NAN", "MBway", "Nantes", 963.58, .1238, .0146, 1.2476, 1.1847, 1.2551, 1, 31533),
       ("MBWAY_PAR", "MBway", "Paris", 1475.02, .1107, .0218, .8150, 1.0597, .8428, 1, 68291),
       ("PIGIER_BOR", "Pigier", "Bordeaux", 1368.09, .0875, .0194, .8787, .8377, .9454, 1, 21968),
       ("PIGIER_LYO", "Pigier", "Lyon", 1704.98, .0850, .0233, .7051, .8135, .7883, 1, 36612),
       ("TUNON_LYO", "Tunon", "Lyon", 1624.68, .0726, .0218, .7399, .6943, .8401, 1, 18850),
       ("TUNON_PAR", "Tunon", "Paris", 2192.89, .0695, .0278, .5482, .6653, .6586, 1, 30240)]
SYN = {"MBWAY_PAR": (398.78, 3292150.80, 8255.51, .1359, 1096393.82, .3330, 2749.35),
       "MBWAY_LYO": (338.50, 2620797.34, 7742.43, .1082, 876442.71, .3344, 2589.21),
       "MBWAY_NAN": (307.34, 2267018.63, 7376.18, .0936, 763702.73, .3369, 2484.86),
       "MBWAY_BOR": (260.41, 1863775.69, 7157.03, .0769, 625402.30, .3356, 2401.59),
       "ISCOM_PAR": (369.35, 3048653.48, 8254.05, .1258, 985765.87, .3233, 2668.90),
       "ISCOM_LIL": (255.09, 1844124.54, 7229.40, .0761, 602778.74, .3269, 2363.04),
       "ISCOM_TLS": (241.48, 1709754.85, 7080.28, .0706, 560009.03, .3275, 2319.06),
       "IPAC_NAN": (151.32, 1067200.79, 7052.49, .0440, 331184.51, .3103, 2188.60),
       "IPAC_REN": (122.12, 818371.29, 6701.28, .0338, 254316.41, .3108, 2082.49),
       "IPAC_MTP": (122.13, 818445.73, 6701.29, .0338, 254308.30, .3107, 2082.23),
       "PIGIER_LYO": (236.95, 1653869.91, 6979.90, .0683, 626670.52, .3789, 2644.77),
       "PIGIER_BOR": (184.72, 1191512.54, 6450.49, .0492, 455099.20, .3820, 2463.77),
       "TUNON_PAR": (143.56, 1133628.58, 7896.53, .0468, 220510.28, .1945, 1536.01),
       "TUNON_LYO": (121.86, 902399.37, 7405.22, .0372, 178759.19, .1981, 1466.92)}

titre_section(p, 16, 2, 12, "CAPACITÉ PAR CAMPUS")
entete(p, 18, 1, ["", "Marque", "Ville", "CAC marginal", "Croissance leads",
                  "Intensité marketing", "Cap effectif", "Cap momentum", "Cap potentiel",
                  "Cap retenu", "Budget acquisition", "Code campus"],
       ["c", "g", "g", "d", "d", "d", "c", "c", "c", "c", "d", "g"])
p.cell(18, 10).fill = fill(JAUNE)
for c in range(1, 13):
    p.cell(18, c).alignment = Alignment(horizontal=p.cell(18, c).alignment.horizontal,
                                        vertical="center", wrap_text=True)

for i, l in enumerate(CAP):
    r = 19 + i
    mettre(p, r, 1, "", F(9), Ce)
    mettre(p, r, 2, l[1], F(9), ind(1))
    mettre(p, r, 3, l[2], F(9), ind(0))
    mettre(p, r, 4, l[3], F(9), Dr, EUR)
    mettre(p, r, 5, l[4], F(9), Dr, PCT)
    mettre(p, r, 6, l[5], F(9), Dr, PCT)
    for j in (6, 7, 8):
        mettre(p, r, j + 1, l[j], F(9), Ce, COEF)
    mettre(p, r, 10, l[9], F(9, True), Ce, COEF, fl=JAUNE)
    mettre(p, r, 11, l[10], F(9), Dr, EUR)
    mettre(p, r, 12, l[0], F(9, False, GRIS), ind(0))
    for c in range(1, 13):
        p.cell(r, c).border = Border(bottom=sd(FIN), left=sd(FIN), right=sd(FIN))

titre_section(p, 36, 2, 12, "SYNTHÈSE PAR CAMPUS")
entete(p, 38, 1, ["", "Marque", "Ville", "Effectif", "Chiffre d'affaires", "Prix moyen",
                  "Part du CA", "EBITDA", "Marge", "EBITDA par étudiant",
                  "Budget acquisition", "CAC marginal"],
       ["c", "g", "g", "d", "d", "d", "d", "d", "c", "d", "d", "d"])
for c in range(1, 13):
    p.cell(38, c).alignment = Alignment(horizontal=p.cell(38, c).alignment.horizontal,
                                        vertical="center", wrap_text=True)
for i, l in enumerate(CAP):
    r = 39 + i
    eff, ca, prix, part, ebitda, marge, epe = SYN[l[0]]
    mettre(p, r, 1, "", F(9), Ce)
    mettre(p, r, 2, l[1], F(9), ind(1))
    mettre(p, r, 3, l[2], F(9), ind(0))
    mettre(p, r, 4, eff, F(9), Dr, NBD)
    mettre(p, r, 5, ca, F(9), Dr, EUR)
    mettre(p, r, 6, prix, F(9), Dr, EUR)
    mettre(p, r, 7, part, F(9), Dr, PCT)
    mettre(p, r, 8, ebitda, F(9), Dr, EUR)
    mettre(p, r, 9, marge, F(9), Ce, PCT)
    mettre(p, r, 10, epe, F(9), Dr, EUR)
    mettre(p, r, 11, l[10], F(9), Dr, EUR)
    mettre(p, r, 12, l[3], F(9), Dr, EUR)
    for c in range(1, 13):
        p.cell(r, c).border = Border(bottom=sd(FIN), left=sd(FIN), right=sd(FIN))

TOT = [(53, "Total campus", 3253.62, 24231703.55, 7447.62, 1.0, 7831343.62, .3232,
        2406.97, 434174, CLAIR, False),
       (54, "Siège et holding", None, None, None, None, -3719839.98, None, None, None,
        "FFFFFF", False),
       (55, "Groupe", 3253.62, 24231703.55, 7447.62, 1.0, 4111503.64, .1697, 1263.67,
        434174, BLEU, True)]
for (r, lib, eff, ca, prix, part, ebitda, marge, epe, bud, fond, dbl) in TOT:
    ligne_total(p, r, 1, 12, fond, dbl)
    mettre(p, r, 2, lib, F(9.5, True), ind(1))
    for c, v, fmt in ((4, eff, NBD), (5, ca, EUR), (6, prix, EUR), (7, part, PCT),
                      (8, ebitda, EUR), (9, marge, PCT), (10, epe, EUR), (11, bud, EUR)):
        if v is not None:
            mettre(p, r, c, v, F(9.5, True), Ce if c == 9 else Dr, fmt)
    ligne_total(p, r, 1, 12, fond, dbl)
    for c in range(1, 13):
        p.cell(r, c).font = F(9.5, True) if p.cell(r, c).value not in (None, "") else F(9)

mettre(p, 57, 2, "Source : Tagetik TGK_MSSQL_07, scénario 2027BUD_V1.", F(8, False, GRIS), ind(0))
p.freeze_panes = "B19"
impression(p, "1:3", "A1:L57")

# ============================================================================
#  NORMALISATION ET CONTROLES
# ============================================================================
for feuille, nl, nc in ((ws, 46, 20), (p, 58, NC2)):
    for r in range(1, nl + 1):
        for c in range(1, nc + 1):
            x = feuille.cell(r, c)
            if x.font is None or x.font.name != UI:
                x.font = F(9)
            if not x.number_format:
                x.number_format = "General"

wb.save(OUT)
v = openpyxl.load_workbook(OUT)
pb = []
for s in v.worksheets:
    if s.merged_cells.ranges:
        pb.append("%s : %d fusion(s)" % (s.title, len(s.merged_cells.ranges)))
    for row in s.iter_rows():
        for c in row:
            if c.number_format == "":
                pb.append("%s!%s : format vide" % (s.title, c.coordinate))
            if c.font and c.font.name not in (None, UI):
                pb.append("%s!%s : police %s" % (s.title, c.coordinate, c.font.name))
print("onglets  :", v.sheetnames)
print("fusions  :", sum(len(s.merged_cells.ranges) for s in v.worksheets))
print("controle :", "PASS" if not pb else pb[:5])

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""simulation_graphes.py — la preuve que les trois query alimentent les graphes.

Rejoue exactement ce que fera la navigation :
  1. execute la logique des trois query reduites sur la donnee reelle
  2. depose chaque resultat la ou Tagetik le deposera, colonnes contigues
  3. branche les trois graphes sur ces plages, en references internes
  4. controle, cellule par cellule, que rien n'est vide et que tout reconcilie

OU TOMBE CHAQUE QUERY (une colonne libre entre deux blocs)

  Q_G1_PONT     AB..AF   ETAPE, SOCLE, ANCRE, HAUSSE, BAISSE
  Q_G2_MARGE    AH..AK   LIBELLE, MARGE_2024, MARGE_2025, MARGE_2026
  Q_G3_TENSION  AM..AO   EXERCICE, IND_DEPENSES, IND_INSCRITS
"""
import csv, sys, openpyxl
from collections import defaultdict
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter as GL
from socle_reel import construire

OUT="NAVIGATION_ALL.xlsm"; N,P=2026,2025
G1,G2,G3 = 28,34,39                      # AB, AH, AM
ORDRE=["Ipac Bachelor Factory","ISCOM","MBway","Pigier","Tunon"]
CODE={"Ipac Bachelor Factory":"IPAC","ISCOM":"ISCOM","MBway":"MBWAY","Pigier":"PIGIER","Tunon":"TUNON"}
LIB={"IPAC_MTP":"Ipac Bachelor Factory Montpellier","IPAC_NAN":"Ipac Bachelor Factory Nantes",
 "IPAC_REN":"Ipac Bachelor Factory Rennes","ISCOM_LIL":"ISCOM Lille","ISCOM_PAR":"ISCOM Paris",
 "ISCOM_TLS":"ISCOM Toulouse","MBWAY_BOR":"MBway Bordeaux","MBWAY_LYO":"MBway Lyon",
 "MBWAY_NAN":"MBway Nantes","MBWAY_PAR":"MBway Paris","PIGIER_BOR":"Pigier Bordeaux",
 "PIGIER_LYO":"Pigier Lyon","TUNON_LYO":"Tunon Lyon","TUNON_PAR":"Tunon Paris"}
C=construire()
acq=defaultdict(lambda: defaultdict(float))
for r in csv.DictReader(open("data/socle_crm.csv",encoding="utf-8-sig"),delimiter=";"):
    acq[r["ENTITY"]][int(r["EXERCICE"])]+=float(r["DEPENSE_ACQ"].replace(",","."))
S=lambda g,a,e: sum(x[a][e] for x in g)

# ------------------------------------------------- Q_G1_PONT, ce qu'elle rend
p_,n_=S(C,"eb",P),S(C,"eb",N)
v =sum((c["eff"][N]-c["eff"][P])*(c["ca"][P]/c["eff"][P]-c["cvar"][P]/c["eff"][P]) for c in C)
pr=sum((c["ca"][N]/c["eff"][N]-c["ca"][P]/c["eff"][P])*c["eff"][N] for c in C)
co=n_-p_-v-pr
R1=[("EBITDA 2025",0,round(p_),0,0),
    ("Activite",   round(p_ if v>0 else p_+v),0,round(max(v,0)),round(max(-v,0))),
    ("Prix / mix", round(p_+v+pr if pr<0 else p_+v),0,round(max(pr,0)),round(max(-pr,0))),
    ("Couts",      round(n_ if co<0 else p_+v+pr),0,round(max(co,0)),round(max(-co,0))),
    ("EBITDA 2026",0,round(n_),0,0)]

# ------------------------------------------------ Q_G2_MARGE, ce qu'elle rend
plusieurs = len({c["marque"] for c in C})>1
grp=defaultdict(list)
for c in C: grp[c["marque"] if plusieurs else c["ent"]].append(c)
R2=[]
for k in sorted(grp,key=lambda x:(ORDRE.index(x) if plusieurs else x)):
    g=grp[k]
    R2.append((k if plusieurs else LIB[k],
               round(S(g,"eb",2024)/S(g,"ca",2024),4),
               round(S(g,"eb",P)/S(g,"ca",P),4),
               round(S(g,"eb",N)/S(g,"ca",N),4)))

# ---------------------------------------------- Q_G3_TENSION, ce qu'elle rend
EX=[2024,P,N]
d={e:sum(acq[c["ent"]][e] for c in C) for e in EX}
i={e:S(C,"inscrits",e) for e in EX}
R3=[(str(e),round(100*d[e]/d[2024],1),round(100*i[e]/i[2024],1)) for e in EX]

# ============================================================ le classeur
wb=openpyxl.load_workbook("COCKPIT_DESIGN.xlsm",keep_vba=True); ws=wb["2"]
for cle in list(wb.defined_names): del wb.defined_names[cle]   # rien qui pointe dehors
ws._charts=[]
for r in range(6,14):                       # on nettoie toute l'ancienne zone
    for c in range(28,60): ws.cell(r,c).value=None

# ------------------------------------------------------------- le tableau
FEUILLE=lambda c: {"D":c["ca"][N],"F":c["eb"][N],"K":c["inscrits"][N],"N":c["eff"][N],
    "O":c["places"][N],"P":c["mix_alt"][N]*c["eff"][N],"Q":c["eb"][P],"R":c["ca"][P],
    "S":acq[c["ent"]][N],"T":acq[c["ent"]][P],"V":c["places"][P],"W":c["inscrits"][P],
    "X":c["eff"][P]}
EBG=S(C,"eb",N)
calc=lambda r: {"E":'=IFERROR((D{0}-R{0})/R{0},"")'.format(r),"G":'=IFERROR((F{0}-Q{0})/Q{0},"")'.format(r),
                "I":'=IFERROR(F{0}/D{0},"")'.format(r),"J":'=IFERROR((I{0}-U{0})*100,"")'.format(r),
                "L":'=IFERROR(N{0}/O{0},"")'.format(r),"M":'=IFERROR(P{0}/N{0},"")'.format(r),
                "U":'=IFERROR(Q{0}/R{0},"")'.format(r),"Y":'=IFERROR((K{0}-W{0})/W{0},"")'.format(r)}
ci=openpyxl.utils.column_index_from_string
lignes=[]; r=38
lignes.append([r,2,"EDUSERVICES",None]); racine=r; r+=1
for m in ORDRE:
    g=[c for c in C if c["marque"]==m]
    tete=r; lignes.append([r,3,m,None]); r+=1
    for c in sorted(g,key=lambda x:LIB[x["ent"]]): lignes.append([r,4,LIB[c["ent"]],c]); r+=1
    lignes[[k for k,l in enumerate(lignes) if l[0]==tete][0]][3]=(tete+1,r-1)
lignes[0][3]=(racine+1,r-1)
for ligne,niv,lib,src in lignes:
    ws.cell(ligne,2,niv); ws.cell(ligne,3,lib)
    if niv==4:
        v2=FEUILLE(src); v2["H"]=src["eb"][N]/EBG
        for k,x in v2.items(): ws.cell(ligne,ci(k),x)
    else:
        a,b=src
        for k in "DFHKNOPQRSTVWX": ws.cell(ligne,ci(k),"=SUBTOTAL(9,{0}{1}:{0}{2})".format(k,a,b))
    for k,f in calc(ligne).items(): ws.cell(ligne,ci(k),f)
DERNIERE=r-1
for rr in range(DERNIERE+1,121):
    for c in range(2,26): ws.cell(rr,c).value=None
ws.cell(5,4,"Forecast 2026"); ws.cell(5,9,"V_FINAL"); ws.cell(5,12,"EDUSERVICES")

# ------------------------------------- les trois resultats de query, en place
ENT={G1:("ETAPE","SOCLE","ANCRE","HAUSSE","BAISSE"),
     G2:("LIBELLE","MARGE_2024","MARGE_2025","MARGE_2026"),
     G3:("EXERCICE","IND_DEPENSES","IND_INSCRITS")}
for c0,titres in ENT.items():
    for j,t in enumerate(titres): ws.cell(6,c0+j,t)
for i2,l in enumerate(R1,7):
    for j,x in enumerate(l): ws.cell(i2,G1+j,x)
for i2,l in enumerate(R2,7):
    for j,x in enumerate(l): ws.cell(i2,G2+j,x)
for i2,l in enumerate(R3,7):
    for j,x in enumerate(l): ws.cell(i2,G3+j,x)
for c in range(16,60): ws.column_dimensions[GL(c)].hidden=True

# --------------------------------------------------------------- les graphes
# References INTERNES uniquement : '2'!$AB$7:$AB$11. Aucun nom defini, aucun
# nom de classeur. C'est ce qui faisait disparaitre les series a chaque
# enregistrement sous un autre nom.
SLATE="526071"; GOOD="1E9E89"; CRIT="D64545"
BLUE="2A78D6"; BLUE2="6FA5DC"; BLUE3="B8CFEC"; ORANGE="F07B32"
def cat(ch,col,n):
    p="'2'!$%s$7:$%s$%d"%(GL(col),GL(col),6+n)
    for x in ch.series: x.cat=AxDataSource(strRef=StrRef(f=p))
    return p
def noms(ch,l):
    for x,t in zip(ch.series,l): x.tx=SeriesLabel(v=t)
def fini(ch):
    ch.height=9.0; ch.width=8.1; ch.visible_cells_only=False
    ch.x_axis.delete=False; ch.y_axis.delete=False
    ch.x_axis.majorTickMark="none"; ch.y_axis.majorTickMark="none"

br=BarChart(); br.type="col"; br.grouping="stacked"; br.overlap=100; br.gapWidth=55
for c in range(G1+1,G1+5):
    br.add_data(Reference(ws,min_col=c,max_col=c,min_row=7,max_row=6+len(R1)),titles_from_data=False)
for x,coul in zip(br.series,(None,SLATE,GOOD,CRIT)):
    if coul is None: x.graphicalProperties.noFill=True
    else: x.graphicalProperties.solidFill=coul; x.graphicalProperties.line.noFill=True
P1=cat(br,G1,len(R1)); br.legend=None; br.y_axis.numFmt='0.0,," M€"'; fini(br)
ws.add_chart(br,"D13")

mg=BarChart(); mg.type="col"; mg.grouping="clustered"; mg.gapWidth=60; mg.overlap=-10
for c in range(G2+1,G2+4):
    mg.add_data(Reference(ws,min_col=c,max_col=c,min_row=7,max_row=6+len(R2)),titles_from_data=False)
for x,coul in zip(mg.series,(BLUE3,BLUE2,BLUE)):
    x.graphicalProperties.solidFill=coul; x.graphicalProperties.line.noFill=True
P2=cat(mg,G2,len(R2)); noms(mg,("2024","2025","2026"))
mg.legend.position="b"; mg.y_axis.numFmt='0%'; fini(mg)
ws.add_chart(mg,"H13")

tn=LineChart()
for c in range(G3+1,G3+3):
    tn.add_data(Reference(ws,min_col=c,max_col=c,min_row=7,max_row=6+len(R3)),titles_from_data=False)
for x,coul in zip(tn.series,(ORANGE,BLUE)):
    x.graphicalProperties.line.solidFill=coul; x.graphicalProperties.line.width=25000
    x.marker=Marker(symbol="circle",size=6); x.smooth=False
    x.marker.graphicalProperties.solidFill=coul; x.marker.graphicalProperties.line.solidFill=coul
P3=cat(tn,G3,len(R3)); noms(tn,("Dépenses","Inscrits"))
tn.legend.position="b"; tn.y_axis.numFmt='0'
tn.y_axis.scaling.min=95; tn.y_axis.scaling.max=125; tn.y_axis.majorUnit=10; fini(tn)
ws.add_chart(tn,"L13")
wb.save(OUT)

# ============================================================== LA PREUVE
import zipfile,re
z=zipfile.ZipFile(OUT)
sh=[n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith("2.xml")][0]
eur=lambda x: f"{x:,.0f}".replace(","," ")
print("="*78); print("  PREUVE — trois query, trois graphes"); print("="*78)
for nom,c0,R,titres in (("Q_G1_PONT",G1,R1,ENT[G1]),("Q_G2_MARGE",G2,R2,ENT[G2]),
                        ("Q_G3_TENSION",G3,R3,ENT[G3])):
    print("\n  %s  ->  %s%d:%s%d   %d colonnes x %d lignes"
          %(nom,GL(c0),7,GL(c0+len(titres)-1),6+len(R),len(titres),len(R)))
    print("    "+"".join("%-15s"%t for t in titres))
    for l in R: print("    "+"".join("%-15s"%(eur(x) if isinstance(x,(int,float)) and abs(x)>999 else x) for x in l))
    vides=[(GL(c0+j),7+i) for i in range(len(R)) for j in range(len(titres))
           if ws.cell(7+i,c0+j).value in (None,"")]
    print("    cellules vides : %s"%(vides or "AUCUNE"))
print("\n"+"-"*78)
print("  CONTROLES")
print("-"*78)
print("  le pont boucle          : %s + %s %s %s = %s   (EBITDA 2026 = %s)"
      %(eur(p_),eur(v),"+" if pr>=0 else "-",eur(abs(pr))+" "+("+" if co>=0 else "-")+" "+eur(abs(co)),
        eur(p_+v+pr+co),eur(n_)))
print("  ecart                   : %.6f EUR"%(p_+v+pr+co-n_))
print("  marges recalculees      : %s"%", ".join("%s %.1f%%"%(l[0][:6],100*l[3]) for l in R2))
print("  indices base 100        : depenses %s | inscrits %s"
      %(" ".join(str(l[1]) for l in R3)," ".join(str(l[2]) for l in R3)))
x=z.read(sh).decode()
print("\n  noms definis dans le classeur : %s"
      %(re.findall(r'<definedName name="([^"]+)"',z.read("xl/workbook.xml").decode()) or "AUCUN"))
for i2,(n2,att) in enumerate([("chart1",P1),("chart2",P2),("chart3",P3)],1):
    c=z.read("xl/charts/chart%d.xml"%i2).decode()
    refs=sorted(set(re.findall(r"<f>([^<]+)</f>",c)))
    print("  chart%d : %d series | categories %s | reference externe : %s"
          %(i2,c.count("<ser>"),att,any(".xls" in r for r in refs)))
    print("           %s"%refs)

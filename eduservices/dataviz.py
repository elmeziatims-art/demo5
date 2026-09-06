#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Cockpit dataviz natif Excel (v2) sur cad + dropdown scenario C5.
Positions v2: trajectoire header 43 / data 44-47 ; reconciliation EBITDA en 13."""
import warnings;warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.chart import LineChart,BarChart,Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.axis import ChartLines
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
wb=openpyxl.load_workbook("MOTEUR_adapte.xlsx")
cad=wb["cad"]
NAVY="1F3864";BLUE="2E86DE";GREEN="27AE60";GOLD="C9971B";ORANGE="E8743B";PURPLE="8E44AD";GREY="9AA5B1"
MARQ=[("MBWAY",BLUE),("ISCOM",GREEN),("IPAC",GOLD),("PIGIER",ORANGE),("TUNON",PURPLE)]
def hdr(ref,txt): cad[ref]=txt; cad[ref].font=Font(size=9,bold=True,color=NAVY)

# ---- dropdown scenario en C5 ----
dv=DataValidation(type="list",formula1='"Cadrage,Optimiste,Prudent"',allow_blank=True)
cad.add_data_validation(dv); dv.add(cad["C5"])

# ---- donnees sources (formules live) ----
hdr("N43","Poids des marques"); cad["N44"]="Marque";cad["O44"]="CA";cad["P44"]="Marge"
for c in ("N44","O44","P44"): cad[c].font=Font(size=8,bold=True,color=GREY)
for i,(m,_) in enumerate(MARQ):
    r=45+i
    cad["N%d"%r]=m
    cad["O%d"%r]='=SUMIFS(Allocation!$K:$K,Allocation!$E:$E,"%s",Allocation!$C:$C,"2026")'%m
    cad["P%d"%r]='=SUMIFS(Allocation!$T:$T,Allocation!$E:$E,"%s",Allocation!$C:$C,"2026")'%m
    cad["O%d"%r].number_format='# ##0';cad["P%d"%r].number_format='# ##0'
hdr("N52","Waterfall EBITDA (cible -> construit)")
cad["N53"]="Etape";cad["O53"]="Base";cad["P53"]="Valeur"
cad["N54"]="Cible";     cad["O54"]=0;               cad["P54"]="=D13"
cad["N55"]="Ecart";     cad["O55"]="=MIN(D13,E13)"; cad["P55"]="=ABS(E13-D13)"
cad["N56"]="Construit"; cad["O56"]=0;               cad["P56"]="=E13"
for r in (54,55,56): cad["O%d"%r].number_format='# ##0';cad["P%d"%r].number_format='# ##0'

def axes(ch):
    ch.x_axis.delete=False;ch.y_axis.delete=False;ch.x_axis.majorGridlines=None
    ch.y_axis.majorGridlines=ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="EEEEEE")))
    ch.x_axis.spPr=GraphicalProperties(ln=LineProperties(solidFill="CCCCCC"))

# 1) TRAJECTOIRE (header 43, data 44-47)
bar=BarChart();bar.type="col";bar.grouping="clustered";bar.title="Trajectoire 2024 -> 2027";bar.height=7.2;bar.width=13
bar.add_data(Reference(cad,min_col=3,max_col=4,min_row=43,max_row=47),titles_from_data=True)
bar.set_categories(Reference(cad,min_col=2,min_row=44,max_row=47))
bar.series[0].graphicalProperties=GraphicalProperties(solidFill=BLUE)
bar.series[1].graphicalProperties=GraphicalProperties(solidFill=GREEN)
bar.y_axis.numFmt='0.0';bar.y_axis.title="M EUR"
line=LineChart();line.add_data(Reference(cad,min_col=5,min_row=43,max_row=47),titles_from_data=True)
line.series[0].graphicalProperties=GraphicalProperties(ln=LineProperties(solidFill=GOLD,w=28000))
line.series[0].marker=Marker(symbol="circle",size=6);line.y_axis.axId=200;line.y_axis.title="Effectif";line.y_axis.crosses="max"
bar+=line;bar.legend.position='b';axes(bar);cad.add_chart(bar,"B49")

# 2) POIDS DES MARQUES (O = CA, rows 44-49)
pm=BarChart();pm.type="bar";pm.title="Poids des marques - CA 2026";pm.height=7.2;pm.width=9;pm.legend=None
pm.add_data(Reference(cad,min_col=15,min_row=44,max_row=49),titles_from_data=True)
pm.set_categories(Reference(cad,min_col=14,min_row=45,max_row=49))
pm.series[0].data_points=[DataPoint(idx=i,spPr=GraphicalProperties(solidFill=MARQ[i][1])) for i in range(5)]
pm.dataLabels=DataLabelList();pm.dataLabels.showVal=True;pm.dataLabels.numFmt='# ##0';pm.x_axis.numFmt='# ##0';axes(pm)
cad.add_chart(pm,"H49")

# 3) WATERFALL (base 53-56 col O, valeur col P)
wf=BarChart();wf.type="col";wf.grouping="stacked";wf.overlap=100;wf.title="Waterfall EBITDA : cible -> construit";wf.height=7.2;wf.width=9;wf.gapWidth=40
wf.add_data(Reference(cad,min_col=15,min_row=53,max_row=56),titles_from_data=True)
wf.add_data(Reference(cad,min_col=16,min_row=53,max_row=56),titles_from_data=True)
wf.set_categories(Reference(cad,min_col=14,min_row=54,max_row=56))
wf.series[0].graphicalProperties=GraphicalProperties(noFill=True)
wf.series[1].data_points=[DataPoint(idx=0,spPr=GraphicalProperties(solidFill=BLUE)),
                          DataPoint(idx=1,spPr=GraphicalProperties(solidFill=GOLD)),
                          DataPoint(idx=2,spPr=GraphicalProperties(solidFill=GREEN))]
wf.legend=None;wf.y_axis.numFmt='# ##0';axes(wf);cad.add_chart(wf,"B64")

# 4) CAP par campus (Pilotage M cap retenu, C ville, rows 12-26)
pil=wb["Pilotage"]
cap=BarChart();cap.type="bar";cap.title="Cap strategique par campus - Cap retenu";cap.height=9;cap.width=9;cap.legend=None
cap.add_data(Reference(pil,min_col=13,min_row=12,max_row=26),titles_from_data=True)
cap.set_categories(Reference(pil,min_col=3,min_row=13,max_row=26))
cap.series[0].graphicalProperties=GraphicalProperties(solidFill=NAVY)
cap.dataLabels=DataLabelList();cap.dataLabels.showVal=True;cap.dataLabels.numFmt='0.00';axes(cap)
cad.add_chart(cap,"H64")

# retirer anciens graphes cad mal places (garder mes 4)
cad._charts=cad._charts[-4:]
wb.save("MOTEUR_adapte.xlsx")
print("OK dataviz v2 + dropdown C5.")
